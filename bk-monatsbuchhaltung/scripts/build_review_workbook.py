from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from datev_io import (
    BATCH_KIND_BOOKING,
    BATCH_KIND_CLARIFICATION,
    STANDARD_BATCH_TYPE,
    accrual_document,
    batch_file_name,
)


NAVY = "183B56"
BLUE = "2F75B5"
PALE_BLUE = "D9EAF7"
GREEN = "C6E0B4"
RED = "F4CCCC"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D6DEE5")
TRAFFIC_COLORS = {"Grün": GREEN, "Rot": RED}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--qa-dir", type=Path)
    return parser.parse_args()


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def amount(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"\s+", "", str(value))
    if not text:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def document_field_1(doc: dict[str, Any]) -> str:
    for line in doc.get("bookings", []):
        if clean(line.get("document_field_1")):
            return clean(line["document_field_1"])
    return clean(doc.get("invoice_number"))


BATCH_KINDS_BY_LIGHT = {
    "Grün": BATCH_KIND_BOOKING,
    "Rot": BATCH_KIND_CLARIFICATION,
}


class BatchFiles:
    """Actual EXTF file names per transaction and booking line from the generator's booking_trace."""

    def __init__(self, data: dict[str, Any]) -> None:
        self.run = data.get("run", {})
        self.by_transaction: dict[str, list[str]] = {}
        self.by_line: dict[tuple[str, int], str] = {}
        for item in data.get("booking_trace", []) or []:
            tid = str(item.get("transaction_id", ""))
            file_name = str(item.get("file", ""))
            if not tid or not file_name:
                continue
            files = self.by_transaction.setdefault(tid, [])
            if file_name not in files:
                files.append(file_name)
            self.by_line[(tid, int(item.get("line", 0) or 0))] = file_name

    def _fallback(self, doc: dict[str, Any]) -> str:
        kind = BATCH_KINDS_BY_LIGHT.get(clean(doc.get("traffic_light")))
        period = clean(doc.get("period")) or clean(self.run.get("buchungsmonat"))
        if not kind or not period:
            return "Nicht bestimmbar"
        batch_type = clean(doc.get("batch_type")) or STANDARD_BATCH_TYPE
        try:
            return batch_file_name(period, kind, batch_type)
        except ValueError:
            return "Nicht bestimmbar"

    def for_document(self, doc: dict[str, Any]) -> str:
        if doc.get("processing_status") != "Buchungszeile erzeugt":
            return "Kein Buchungsstapel"
        files = self.by_transaction.get(str(doc.get("transaction_id", "")))
        return ", ".join(sorted(files)) if files else self._fallback(doc)

    def for_line(self, doc: dict[str, Any], line_no: int) -> str:
        if doc.get("processing_status") != "Buchungszeile erzeugt":
            return "Kein Buchungsstapel"
        return self.by_line.get((str(doc.get("transaction_id", "")), line_no)) or self.for_document(doc)


def compact_posting(doc: dict[str, Any]) -> str:
    if doc.get("processing_status") != "Buchungszeile erzeugt":
        return clean(doc.get("processing_status")) or "Keine Buchungszeile"
    rows: list[str] = []
    for line in doc.get("bookings", []):
        value = amount(line.get("amount"))
        value_text = (
            f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            if value is not None else clean(line.get("amount"))
        )
        account = " ".join(filter(None, [
            clean(line.get("account")), clean(line.get("account_name"))
        ]))
        contra = " ".join(filter(None, [
            clean(line.get("contra_account")), clean(line.get("contra_account_name"))
        ]))
        rows.append(
            f"{value_text} {doc.get('currency', 'EUR')} {line.get('debit_credit', '')}"
            f" | {account} an {contra} | BU {line.get('bu_key') or '–'}"
            f" | BF1 {line.get('document_field_1') or document_field_1(doc)}"
        )
    return "\n".join(rows)


def apply_header(ws, row: int, start: int, end: int) -> None:
    for cell in ws.iter_cols(min_col=start, max_col=end, min_row=row, max_row=row):
        item = cell[0]
        item.fill = PatternFill("solid", fgColor=BLUE)
        item.font = Font(color=WHITE, bold=True)
        item.alignment = Alignment(vertical="center", wrap_text=True)
        item.border = Border(bottom=THIN)


def set_widths(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width


def traffic_format(ws, end_row: int) -> None:
    if end_row < 2:
        return
    for row in range(2, end_row + 1):
        cell = ws.cell(row, 1)
        color = TRAFFIC_COLORS.get(clean(cell.value))
        if color:
            # Als echte Zellformatierung setzen, damit die Ampelfarbe auch in
            # Vorschauen ohne Auswertung bedingter Formatierung sichtbar ist.
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True)

    # Bedingte Formatierung hält die Farbe bei späteren Änderungen korrekt.
    for value, color in TRAFFIC_COLORS.items():
        ws.conditional_formatting.add(
            f"A2:A{end_row}",
            FormulaRule(
                formula=[f'$A2="{value}"'],
                fill=PatternFill("solid", fgColor=color),
                font=Font(bold=True),
            ),
        )


def add_title(ws, title: str, end_column: int) -> None:
    ws.insert_rows(1, 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.sheet_view.showGridLines = False


def build(data: dict[str, Any], output: Path) -> None:
    run = data["run"]
    batch_files = BatchFiles(data)
    cost_config = run.get("cost_center_config") if isinstance(run.get("cost_center_config"), dict) else None
    with_kost2 = bool(cost_config and (cost_config.get("kost2_required") is True or cost_config.get("kost2_allowed")))
    order = {"Rot": 0, "Grün": 1, None: 9, "": 9}
    if any(doc.get("traffic_light") not in {"Grün", "Rot", None, ""} for doc in data.get("documents", [])):
        raise ValueError("Nur Grün und Rot sind zulässig.")
    docs = sorted(
        data.get("documents", []),
        key=lambda doc: (
            order.get(doc.get("traffic_light"), 9),
            str(doc.get("period", "")),
            str(doc.get("recognized_date", "")),
            str(doc.get("transaction_id", "")),
        ),
    )
    posting_docs = docs + [accrual_document(release, run, number)
                           for number, release in enumerate(data.get("accrual_releases", []), start=1)]
    cases: dict[str, dict[str, Any]] = {}
    for case in data.get("clarification_cases", []):
        for transaction_id in case.get("transaction_ids", []):
            cases.setdefault(str(transaction_id), case)

    wb = Workbook()
    wb.remove(wb.active)
    guide = wb.create_sheet("Anleitung")
    summary = wb.create_sheet("Übersicht")
    review = wb.create_sheet("Belegprüfung")
    bookings = wb.create_sheet("Buchungszeilen")
    notes = wb.create_sheet("Mandanten-Hinweise")
    masters = wb.create_sheet("Stammdatenänderungen")
    rate_sheet = wb.create_sheet("Klärungsquote")

    guide_rows = [
        ["Schritt / Feld", "Bedeutung"],
        ["1. Belegprüfung", "Rote Fälle bearbeiten. Grüne Vorgänge stehen im Buchungsstapel, rote im Klärungsstapel; beide Stapel werden in DATEV importiert, der Klärungsstapel wird dort bearbeitet und erst danach festgeschrieben. Rechts neben der Ampel steht der vollständige Dateiname des jeweiligen Stapels. Die Spalte Belegdatum laut Beleg zeigt das sicher erkannte Datum; im Klärungsstapel ist das DATEV-Belegdatum immer leer und wird in DATEV nachgetragen."],
        ["2. Buchungszeilen", "Direkt rechts neben der Ampel steht der DATEV-Buchungsstapel; danach Konten, BU-Schlüssel, Belegfeld 1, Buchungstext und Periode nachvollziehen."],
        ["3. Rücklaufstatus", "Für jeden roten Vorgang einen Abschlussstatus wählen: unverändert übernommen, geändert oder nicht übernommen. Offen ist kein Abschlussstatus."],
        ["4. Mitarbeiter-Ergebnis", "Bei geändert oder nicht übernommen ist die endgültige Behandlung als Mitarbeiter-Ergebnis Pflicht."],
        ["Grün", "Vollständig und plausibel; keine offene fachliche Frage."],
        ["Rot", "Aktive Bearbeitung erforderlich. Nur konkret ungeklärte Felder bleiben leer; bei Anlagenzugängen bleibt das Anlagenkonto immer offen. Arbeitsanweisungen stehen ausschließlich hier."],
        ["Zahlungsavise", "Nicht buchen. Das gesonderte Belegtransfer_Avise-ZIP in DATEV Unternehmen online hochladen."],
        ["DATEV-Import", "1. Stammdaten, 2. reguläre Belegtransfer-ZIPs, 3. Avis-ZIPs, 4. Buchungsstapel je Monat (danach konfigurierte Stapeltypen wie _Eigenbelege derselben Periode), 5. Klärungsstapel je Monat als eigener Importvorgang, ebenfalls importieren; erst festschreiben, wenn alle roten Zeilen bearbeitet sind. Alle Stapel, mit und ohne Klärung, werden übertragen. DATEV-Testimportstatus beachten."],
        ["Klärungsquote", "Blatt Klärungsquote zeigt N (buchungsrelevante Vorgänge), R (rote Vorgänge) und Q = 100 × R / N vor und nach der Zweitprüfung, die Grenzstufe, die Verteilung nach Rot-Gründen und die verbleibenden offenen Gründe. Die Quote ist ein Qualitätsindikator, kein Zielwert."],
        ["Technisch nicht auswertbar", "Vorgänge mit diesem Endstatus wurden nach dokumentiertem Auswertungsversuch einschließlich Belegbildprüfung nicht ausgewertet; sie stehen ohne Ampel in der Belegprüfung mit den Versuchen als nächstem Schritt."],
        ["Belegdateien", "Jeder Buchungsbeleg ist im Belegtransfer genau eine eigene PDF-Datei; Sammel-PDFs wurden je Vorgang getrennt, verteilte oder Bildbelege zu einer PDF zusammengeführt."],
    ]
    if cost_config:
        allowed = ", ".join(f"{number} {name}" for number, name in cost_config.get("kost1_allowed", {}).items())
        pflicht = "Pflicht auf jeder Buchungszeile" if run.get("kostenstellenpflicht") is True else "wird gesetzt, wenn ableitbar; sonst leer ohne Ampelwirkung"
        guide_rows.append(["Kostenstellen", f"KOST1 für diesen Mandanten: {pflicht}. Zulässige Kostenstellen: {allowed}. Offenes kost1 steht wie andere offene Felder in der Begründungsspalte."])
    for row in guide_rows:
        guide.append(row)
    apply_header(guide, 1, 1, 2)
    for row in range(2, guide.max_row + 1):
        guide.cell(row, 1).font = Font(bold=True)
        guide.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row].height = 34
    for row, color in ((6, GREEN), (7, RED)):
        guide.cell(row, 1).fill = PatternFill("solid", fgColor=color)
    set_widths(guide, [25, 90])
    guide.freeze_panes = "A2"
    add_title(guide, f"Buchungsprüfung {run['mandantennummer']} – {run['buchungsmonat']}", 8)

    summary.append(["Kennzahl", "Gesamt", "Grün", "Rot"])
    summary.append(["Belege", len(docs)] + [
        sum(1 for doc in docs if doc.get("traffic_light") == value)
        for value in ("Grün", "Rot")
    ])
    summary.append(["Belegsumme", sum(amount(doc.get("total_amount")) or 0 for doc in docs)] + [
        sum((amount(doc.get("total_amount")) or 0) for doc in docs if doc.get("traffic_light") == value)
        for value in ("Grün", "Rot")
    ])
    summary.append(["Buchungszeilen", sum(len(doc.get("bookings", [])) for doc in posting_docs)] + [
        sum(len(doc.get("bookings", [])) for doc in posting_docs if doc.get("traffic_light") == value)
        for value in ("Grün", "Rot")
    ])
    summary.append([])
    summary.append(["Kontrollpunkt", "Ergebnis"])
    batch_rows = data.get("booking_batches") or []
    controls = [
        ("Vorperiodenbelege", sum(1 for doc in docs if str(doc.get("period", "")) < run["buchungsmonat"])),
        ("Zukunftsbelege", sum(1 for doc in docs if str(doc.get("period", "")) > run["buchungsmonat"])),
        ("Sichere Dubletten", sum(1 for doc in docs if doc.get("processing_status") == "sichere Dublette – nicht erneut gebucht")),
        ("Zahlungsavise", sum(1 for doc in docs if doc.get("payment_advice"))),
        ("Nicht buchungsrelevant", sum(1 for doc in docs if doc.get("processing_status") == "nicht buchungsrelevant")),
        ("Technisch nicht auswertbar", sum(1 for doc in docs if doc.get("processing_status") == "technisch nicht auswertbar")),
        ("Klärungsquote nach Zweitprüfung", _quota_text(((data.get("clarification_rate") or {}).get("after") or {}).get("Q", "nicht berechnet"))),
        ("Laufstatus (Abschluss-Gate)", (data.get("run_completion") or {}).get("status", "nicht ausgewiesen")),
        ("Offene Klärfälle", len(data.get("clarification_cases", []))),
        ("Neue/geänderte Stammdaten", len(data.get("master_records", []))),
        ("Mandanten-Hinweise", len(data.get("client_notes", []))),
        ("Vollständigkeitskontrolle", "VOLLSTÄNDIG"),
    ]
    for item in controls:
        summary.append(list(item))
    summary.append([])
    batch_header_row = summary.max_row + 1
    summary.append(["DATEV-Stapel", "Stapelbezeichnung", "Zeilen", "Summe"])
    if batch_rows:
        for batch in batch_rows:
            summary.append([batch.get("file", ""), batch.get("label", ""), int(batch.get("rows", 0)), amount(batch.get("amount_total"))])
    else:
        per_file: dict[str, tuple[int, float]] = {}
        for doc in posting_docs:
            for line_no, line in enumerate(doc.get("bookings", []), start=1):
                name = batch_files.for_line(doc, line_no)
                count, total = per_file.get(name, (0, 0.0))
                per_file[name] = (count + 1, total + (amount(line.get("amount")) or 0.0))
        for name, (count, total) in sorted(per_file.items()):
            summary.append([name, "", count, total])
    apply_header(summary, 1, 1, 4)
    apply_header(summary, 6, 1, 2)
    apply_header(summary, batch_header_row, 1, 4)
    for row in range(2, 5):
        for col in range(2, 5):
            summary.cell(row, col).number_format = '#,##0.00' if row == 3 else '#,##0'
    for row in range(batch_header_row + 1, summary.max_row + 1):
        summary.cell(row, 3).number_format = '#,##0'
        summary.cell(row, 4).number_format = '#,##0.00'
    set_widths(summary, [44, 24, 16, 16])
    summary.freeze_panes = "A2"
    add_title(summary, f"Übersicht {run['mandantennummer']} – {run['buchungsmonat']}", 4)

    review_headers = [
        "Ampel-Einstufung", "Buchungsstapel", "Vorgangs-ID", "Belegdatum laut Beleg", "Geschäftspartner",
        "Belegfeld 1", "Betrag", "Währung", "Buchungsperiode", "Kontierung",
        "Ableitung", "Prüfergebnis / Ampelbegründung",
        "Offener Punkt / nächster Schritt", "Bearbeitungsstatus", "Mitarbeiter-Ergebnis",
    ]
    if cost_config:
        review_headers.insert(review_headers.index("Kontierung") + 1, "KOST1")
    review_col = {name: index for index, name in enumerate(review_headers, start=1)}
    review.append(review_headers)
    for doc in docs:
        case = cases.get(str(doc.get("transaction_id", "")), {})
        next_step = " – ".join(dict.fromkeys(filter(None, [
            clean(case.get("recommendation")), clean(case.get("decision_needed"))
        ])))
        open_details = "; ".join(dict.fromkeys(
            f"{field}: {clean(reason)}" for line in doc.get("bookings", [])
            for field, reason in line.get("open_fields", {}).items()
        ))
        if not next_step:
            if doc.get("payment_advice"):
                next_step = "Gesondertes Avis-ZIP in DATEV Unternehmen online hochladen."
            elif doc.get("processing_status") == "sichere Dublette – nicht erneut gebucht":
                next_step = "Keine erneute Buchung."
            elif doc.get("processing_status") == "nicht buchungsrelevant":
                next_step = "Keine Buchung."
            elif doc.get("processing_status") == "technisch nicht auswertbar":
                attempts = "; ".join(
                    f"{clean(item.get('method'))}: {clean(item.get('result'))}"
                    for item in doc.get("evaluation_attempts", []) or [] if isinstance(item, dict)
                )
                next_step = "Technisch nicht auswertbar nach dokumentiertem Auswertungsversuch (" + attempts + "); Original anfordern."
            elif doc.get("traffic_light") == "Grün":
                next_step = "Keine weitere Bearbeitung."
        row_values = [
            doc.get("traffic_light") or "",
            batch_files.for_document(doc),
            doc.get("transaction_id") or "",
            datetime.strptime(doc["recognized_date"], "%Y-%m-%d") if doc.get("recognized_date") else None,
            doc.get("partner") or "",
            document_field_1(doc),
            amount(doc.get("total_amount")),
            doc.get("currency") or "",
            doc.get("period") or "",
            compact_posting(doc),
            clean(doc.get("derivation")),
            " – ".join(filter(None, [re.sub(r"^(Grün|Rot)\s*:\s*", "", clean(doc.get("reason")), flags=re.I), clean(case.get("booking_risk"))])),
            " – ".join(filter(None, [open_details, next_step])),
            "offen" if doc.get("traffic_light") == "Rot" else "",
            "",
        ]
        if cost_config:
            kost_values = [
                clean(line.get("kost1")) or ("offen" if "kost1" in (line.get("open_fields") or {}) else "")
                for line in doc.get("bookings", [])
            ]
            if doc.get("processing_status") == "Buchungszeile erzeugt" and not any(kost_values):
                kost_text = "ohne Kostenstelle"
            else:
                kost_text = ", ".join(dict.fromkeys(value for value in kost_values if value))
            row_values.insert(review_col["KOST1"] - 1, kost_text)
        review.append(row_values)
    column_count = len(review_headers)
    apply_header(review, 1, 1, column_count)
    review.freeze_panes = "C2"
    review.auto_filter.ref = review.dimensions
    widths = [18, 38, 14, 14, 27, 22, 14, 10, 16, 58, 52, 48, 48, 27, 48]
    if cost_config:
        widths.insert(review_col["KOST1"] - 1, 14)
    set_widths(review, widths)
    status_col = review_col["Bearbeitungsstatus"]
    result_col = review_col["Mitarbeiter-Ergebnis"]
    status_letter = review.cell(1, status_col).column_letter
    result_letter = review.cell(1, result_col).column_letter
    for row in range(2, review.max_row + 1):
        review.cell(row, review_col["Belegdatum laut Beleg"]).number_format = "dd.mm.yyyy"
        review.cell(row, review_col["Betrag"]).number_format = '#,##0.00'
        for col in range(review_col["Kontierung"], column_count + 1):
            review.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
        review.row_dimensions[row].height = 52
    traffic_format(review, review.max_row)
    status_validation = DataValidation(
        type="list",
        formula1='"offen,unverändert übernommen,geändert,nicht übernommen"',
    )
    review.add_data_validation(status_validation)
    status_validation.add(f"{status_letter}2:{status_letter}{max(2, review.max_row)}")
    review.conditional_formatting.add(
        f"{result_letter}2:{result_letter}{max(2, review.max_row)}",
        FormulaRule(
            formula=[f'AND(OR(${status_letter}2="geändert",${status_letter}2="nicht übernommen"),${result_letter}2="")'],
            fill=PatternFill("solid", fgColor=RED),
            font=Font(color="9C0006", bold=True),
        ),
    )

    booking_headers = [
        "Ampel-Einstufung", "Buchungsstapel", "Vorgangs-ID", "Betrag", "Soll/Haben", "Konto",
        "Kontobezeichnung", "Gegenkonto", "Gegenkontobezeichnung",
        "BU-Schlüssel", "Erkanntes Belegdatum", "Belegfeld 1",
        "Buchungstext", "Buchungsperiode",
    ]
    booking_widths = [18, 38, 14, 14, 12, 12, 26, 14, 28, 13, 19, 23, 42, 16]
    if cost_config:
        position = booking_headers.index("BU-Schlüssel") + 1
        booking_headers.insert(position, "KOST1")
        booking_widths.insert(position, 12)
        if with_kost2:
            booking_headers.insert(position + 1, "KOST2")
            booking_widths.insert(position + 1, 12)
    booking_col = {name: index for index, name in enumerate(booking_headers, start=1)}
    bookings.append(booking_headers)
    for doc in posting_docs:
        for line_no, line in enumerate(doc.get("bookings", []), start=1):
            values = [
                doc.get("traffic_light") or "", batch_files.for_line(doc, line_no),
                doc.get("transaction_id") or "", amount(line.get("amount")), line.get("debit_credit") or "",
                str(line.get("account") or ""), line.get("account_name") or "",
                str(line.get("contra_account") or ""), line.get("contra_account_name") or "",
                line.get("bu_key") or "",
                datetime.strptime(doc["recognized_date"], "%Y-%m-%d") if doc.get("recognized_date") else None,
                line.get("document_field_1") or document_field_1(doc),
                line.get("booking_text") or "", doc.get("period") or "",
            ]
            if cost_config:
                values.insert(booking_col["KOST1"] - 1, clean(line.get("kost1")))
                if with_kost2:
                    values.insert(booking_col["KOST2"] - 1, clean(line.get("kost2")))
            bookings.append(values)
    apply_header(bookings, 1, 1, len(booking_headers))
    bookings.freeze_panes = "C2"
    bookings.auto_filter.ref = bookings.dimensions
    set_widths(bookings, booking_widths)
    for row in range(2, bookings.max_row + 1):
        bookings.cell(row, booking_col["Betrag"]).number_format = '#,##0.00'
        bookings.cell(row, booking_col["Erkanntes Belegdatum"]).number_format = "dd.mm.yyyy"
        bookings.cell(row, booking_col["Buchungstext"]).alignment = Alignment(wrap_text=True)
    traffic_format(bookings, bookings.max_row)

    notes.append(["Vorgangs-ID", "Geschäftspartner", "Hinweis", "Empfohlenes Vorgehen"])
    for item in data.get("client_notes", []):
        notes.append([
            item.get("transaction_id") or "", item.get("partner") or "",
            item.get("note") or "", item.get("suggested_action") or "",
        ])
    apply_header(notes, 1, 1, 4)
    notes.freeze_panes = "A2"
    notes.auto_filter.ref = notes.dimensions
    set_widths(notes, [14, 28, 58, 58])
    for row in range(2, notes.max_row + 1):
        notes.cell(row, 3).alignment = Alignment(wrap_text=True)
        notes.cell(row, 4).alignment = Alignment(wrap_text=True)

    masters.append([
        "Aktion", "Konto", "Typ", "Name", "USt-ID", "Bankverbindungen",
        "Vollständiger Datensatz", "Hinweis",
    ])
    for record in data.get("master_records", []):
        banks = "; ".join(filter(None, [
            " / ".join(filter(None, [clean(bank.get("iban")), clean(bank.get("bic")), clean(bank.get("bank_name"))]))
            for bank in record.get("banks", [])
        ]))
        complete = record.get("action") == "Neuanlage" or record.get("full_current_record_available") is True
        masters.append([
            record.get("action") or "", str(record.get("account") or ""),
            record.get("account_type") or "", record.get("name") or "",
            record.get("vat_id") or "", banks, "JA" if complete else "NEIN",
            "" if complete else "Vollständigen aktuellen Datensatz manuell pflegen.",
        ])
    apply_header(masters, 1, 1, 8)
    masters.freeze_panes = "A2"
    masters.auto_filter.ref = masters.dimensions
    set_widths(masters, [16, 14, 14, 34, 20, 60, 22, 50])
    for row in range(2, masters.max_row + 1):
        masters.cell(row, 6).alignment = Alignment(wrap_text=True)
        masters.cell(row, 8).alignment = Alignment(wrap_text=True)

    _build_rate_sheet(rate_sheet, data.get("clarification_rate") or {})

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def _quota_text(value: Any) -> str:
    if value is None:
        return "nicht berechenbar"
    if isinstance(value, (int, float)):
        return f"{value} %"
    return str(value)


def _build_rate_sheet(ws, rate: dict[str, Any]) -> None:
    """Klärungsquoten-Nachweis als Registerblatt (VALIDIERUNG.md, Klärungsquote und Zweitprüfung)."""
    before = rate.get("before") or {}
    after = rate.get("after") or {}
    counts = rate.get("counts") or {}
    ws.append(["Kennzahl", "vor Zweitprüfung", "nach Zweitprüfung"])
    ws.append(["N (buchungsrelevante Vorgänge, einmalig gezählt)", before.get("N", "nicht berechnet"), after.get("N", "nicht berechnet")])
    ws.append(["R (rote Vorgänge, Mehrfachzeilen einmal)", before.get("R", "nicht berechnet"), after.get("R", "nicht berechnet")])
    ws.append(["Q = 100 × R / N", _quota_text(before.get("Q", "nicht berechnet")), _quota_text(after.get("Q", "nicht berechnet"))])
    ws.append(["Grenzstufe", before.get("stage_label", ""), after.get("stage_label", "")])
    ws.append([])
    ws.append(["Prüfung", "Ergebnis"])
    ws.append(["Prüfdatum", rate.get("checked_at", "")])
    ws.append(["Prüfschritt", rate.get("check_step", "")])
    ws.append(["Zweitprüfung erforderlich", "ja" if rate.get("second_review_required") else "nein"])
    ws.append(["Zweitprüfung durchgeführt", "ja" if rate.get("second_review_performed") else "nein"])
    ws.append(["Konkret nachgeprüfte Fälle", rate.get("reviewed_cases", 0)])
    ws.append(["Fachlich auf Grün korrigiert", len(rate.get("corrected_to_green", []))])
    ws.append(["Verbleibende rote Fälle", rate.get("remaining_red", 0)])
    ws.append(["Rote Buchungszeilen (nicht rote Fälle)", counts.get("rote_buchungszeilen", 0)])
    ws.append(["Technisch nicht auswertbar", counts.get("technisch_nicht_auswertbar", 0)])
    ws.append(["Bereits in DATEV vorhanden", counts.get("bereits_in_datev_vorhanden", 0)])
    ws.append(["Sichere Dubletten", counts.get("sichere_dubletten", 0)])
    ws.append(["Aussteuerungen", counts.get("aussteuerungen", 0)])
    ws.append(["Ergebnis", rate.get("result", "nicht berechnet")])
    ws.append([])
    ws.append(["Rot-Grund", "Anzahl", "Ursachen-/Plausibilitätsprüfung"])
    labels = rate.get("reason_labels") or {}
    cause = rate.get("cause_analysis") or {}
    distribution = rate.get("reason_distribution") or {}
    if not distribution:
        ws.append(["–", 0, "keine roten Vorgänge"])
    for code, count in distribution.items():
        ws.append([f"{labels.get(code, code)} ({code})", count, cause.get(code, "")])
    ws.append([])
    ws.append(["Vorgangs-ID", "Rot-Grund", "Offene Felder", "Belegreferenz", "Durchgeführter Prüfversuch", "Nächster Prüfschritt"])
    for row in rate.get("red_cases", []) or []:
        ws.append([
            row.get("transaction_id", ""), f"{row.get('label', '')} ({row.get('code', '')})",
            ", ".join(row.get("open_fields", [])), row.get("document_reference", ""),
            row.get("verification_attempted", ""), row.get("next_check", ""),
        ])
    apply_header(ws, 1, 1, 3)
    apply_header(ws, 7, 1, 2)
    set_widths(ws, [46, 28, 48, 22, 48, 48])
    for row in range(2, ws.max_row + 1):
        for col in range(1, 7):
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def main() -> int:
    args = parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    build(data, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
