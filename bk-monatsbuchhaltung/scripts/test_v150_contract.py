"""Regressionstests Änderungsanweisung v1.5.0 (Übertragung über den Riecken-Connector).

1. Versionstest: Skill 1.5.0, Paketvertrag der Skripte 1.4.1, Abschnitt 5 und Referenzen vorhanden
2. riecken_records.py: grüne Zeilen 1:1, rote Zeilen über das Klärungskonto, KLÄR-Text höchstens 60 Zeichen,
   keine Zeile Klärungskonto an Klärungskonto, kein Platzhalterbetrag, Summenabgleich je Stapel,
   Abweisung bei zwei offenen Kontoseiten, fehlendem Betrag, fehlendem Datum und überlangem Vorgabetext
3. update_review_workbook_riecken.py: Ursprungsfassung unter ersetzt/, Riecken-Stapel in Belegprüfung und
   Buchungszeilen, Klärungskonto und KLÄR-Text, Nutzerentscheidung, Blatt Riecken-Übertragung, Statusdatei,
   EXTF-Dateien verschoben, riecken_transfer im Laufmanifest
4. remove_document_from_transfer.py: Beleg samt document.xml-Eintrag entfernt, testzip, Original unter ersetzt/
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from openpyxl import load_workbook

import build_package
import remove_document_from_transfer
import riecken_records
import validate_package
from test_datev_contract import red_reason
from test_v140_contract import Scenario, make_run, red_booking
from test_v141_contract import build_ok, write_run

SKILL_ROOT = SCRIPT_DIR.parent
SKILL_VERSION = "1.5.0"
CONTRACT_VERSION = "1.4.1"
FORBIDDEN_NEXT = "1.5.1"
CLEARING = "1599"
CLEARING_NAME = "Klärungskonto Buchhaltung"


# 1. Version und Texte --------------------------------------------------------------

def test_version_and_texts() -> None:
    skill = (SKILL_ROOT / "SKILL.md").read_text(encoding="utf-8")
    assert f"# BK Monatsbuchhaltung v{SKILL_VERSION}" in skill
    assert f"Version `{SKILL_VERSION}`" in skill and f"Paketvertrag der Skripte (`build_package.py`, `validate_package.py`) bleibt Version `{CONTRACT_VERSION}`" in skill
    assert f"Startnachweis: bk-monatsbuchhaltung v{SKILL_VERSION} (Paketvertrag {CONTRACT_VERSION})" in skill
    assert "Riecken-Übertragung nur auf ausdrücklichen Auftrag mit Klärungskonto" in skill
    assert "### 5. Übertragung über den Riecken-Connector (nur auf ausdrücklichen Auftrag)" in skill
    assert skill.index("### 5. Übertragung über den Riecken-Connector") < skill.index("## Übergabe und Selbstbegrenzung")
    for needle in (
        "Überträgt auf ausdrücklichen Auftrag Personenkonten, Buchungsstapel und Klärungsposten über den Riecken-Connector nach DATEV.",
        "Eine Übertragung über Riecken erfolgt ausschließlich nach Abschnitt 5 und nur auf ausdrücklichen Auftrag nach der Paketübergabe.",
        "Das Klärungskonto nach Abschnitt 5 gilt ausschließlich für die Übertragung über den Riecken-Connector.",
        "Ausnahme: Klärungsbuchungen nach Abschnitt 5 tragen das Kürzel `KLÄR`",
        "`159900 Klärungskonto Buchhaltung`", "nie Klärungskonto an Klärungskonto", f"Stapelbezeichnung `Rechnungen Nachlauf`", "insbesondere kein 1 Cent",
        "`KLÄR <Vorgangs-ID> <offenes Thema> <Zielkonto oder Alternativen>`", "00_STATUS_NACH_RIECKEN_UEBERTRAGUNG.md",
        "scripts/riecken_records.py", "scripts/update_review_workbook_riecken.py", "scripts/remove_document_from_transfer.py",
    ):
        assert needle in skill, needle
    ui = (SKILL_ROOT / "agents" / "openai.yaml").read_text(encoding="utf-8")
    assert f"BK Monatsbuchhaltung v{SKILL_VERSION}" in ui and "Riecken-Übertragung mit Klärungskonto" in ui
    assert build_package.SKILL_VERSION == CONTRACT_VERSION and validate_package.EXPECTED_SKILL_VERSION == CONTRACT_VERSION
    decisions = (SKILL_ROOT / "references" / "FACHLICHE_ENTSCHEIDUNGEN.md").read_text(encoding="utf-8")
    assert "## 19. Riecken-Übertragung und Klärungskonto" in decisions and "Das EXTF-Paket selbst bleibt ohne Ersatzkonten." in decisions
    validation = (SKILL_ROOT / "references" / "VALIDIERUNG.md").read_text(encoding="utf-8")
    assert "## Kontrolle nach Riecken-Übertragung" in validation and "im EXTF-Paket; für die Riecken-Übertragung siehe `SKILL.md` Abschnitt 5" in validation
    assert "kein EXTF-Testimport nötig" in validation
    schema = (SKILL_ROOT / "references" / "EINGABESCHEMA.md").read_text(encoding="utf-8")
    assert "`riecken_transfer`" in schema and "ausschließlich in `SKILL.md` Abschnitt 5" in schema
    template = (SKILL_ROOT / "assets" / "mandantenprofil_vorlage.md").read_text(encoding="utf-8")
    assert "- Ersatzkonten im EXTF-Paket: nicht verwenden; ungeklärte Felder offen lassen." in template
    assert "- klaerungskonto (nur für Riecken-Übertragung, Arbeitskonto mit Zielsaldo 0): 159900 Klärungskonto Buchhaltung" in template
    for path in [SKILL_ROOT / "SKILL.md", SKILL_ROOT / "agents" / "openai.yaml", *sorted((SKILL_ROOT / "references").glob("*.md"))]:
        assert FORBIDDEN_NEXT not in path.read_text(encoding="utf-8"), f"{path.name} verweist auf {FORBIDDEN_NEXT}"


# 2. Records ------------------------------------------------------------------------

def make_scenario(root: Path) -> Scenario:
    scenario = Scenario(root, make_run(root))
    scenario.run["account_config"]["clarification"] = CLEARING
    scenario.add("Grün")
    scenario.add("Grün", total_amount="50.00", document_type="Gutschrift", bookings=[red_booking(amount="50.00", debit_credit="H", booking_text="Gutschrift DATEV Test", open_fields={})])
    # V0003: Konto offen -> Klärungskonto im Soll
    scenario.add("Rot", reason="Kontierung offen.", bookings=[red_booking(account=None, account_name=None, open_fields={"account": "Kontierung aus Beleg nicht erkennbar."})])
    # V0004: Anlage, Anlagenkonto offen -> Klärungskonto, Zielkonto im Text
    scenario.add("Rot", reason="Anlagenzugang.", asset_booking=True, total_amount="238.00", red_reason=red_reason("anlage_gwg_spezialregel"),
                 bookings=[red_booking(amount="238.00", account=None, account_name=None, asset_account_field="account", open_fields={"account": "Anlagenkonto bleibt offen."})])
    # V0005: beide Kontoseiten offen -> nicht übertragbar
    scenario.add("Rot", reason="Rechtsträger und Konto unklar.", red_reason=red_reason("rechtstraeger_unklar"),
                 bookings=[red_booking(account=None, account_name=None, contra_account=None, contra_account_name=None,
                                       open_fields={"account": "Konto unklar.", "contra_account": "Geschäftspartneridentität unklar."})])
    # V0006: Betrag unsicher -> nicht übertragbar
    scenario.add("Rot", reason="Betrag unsicher, Seite 2 fehlt.", total_amount=None, red_reason=red_reason("fehlende_belegangabe"),
                 bookings=[red_booking(amount=None, open_fields={"amount": "Seite 2 der Rechnung fehlt."})])
    # V0007: Datum unsicher -> nicht übertragbar
    scenario.add("Rot", reason="Belegdatum nicht erkennbar.", recognized_date=None, red_reason=red_reason("fehlende_belegangabe"),
                 bookings=[red_booking(open_fields={"recognized_date": "Datum auf dem Beleg nicht lesbar."})])
    return scenario


def extf_total(path: Path) -> Decimal:
    return sum((riecken_records.extf_amount(row[0]) or Decimal("0")) for row in riecken_records.extf_rows(path))


def test_records(root: Path) -> tuple[Path, Path, dict, list]:
    scenario = make_scenario(root)
    package, manifest, _ = build_ok(scenario, "riecken")
    data = json.loads((scenario.root / "riecken.json").read_text(encoding="utf-8"))
    datev = package / "01_DATEV_Import"

    # Überlanger Vorgabetext wird abgewiesen, nicht gekürzt.
    long_text = "KLÄR V0003 " + "Konto unklar, Beleg ohne Leistungsbeschreibung, Lieferant fragen " * 2
    _, rejected = riecken_records.build_records(package, data, CLEARING, CLEARING_NAME, {"V0003": long_text}, "Eingangsrechnungen", "Rechnungen Nachlauf")
    assert any(item["transaction_id"] == "V0003" and riecken_records.REASON_TEXT in item["reason"] for item in rejected)

    # Klärungskonto muss zu account_config.clarification passen.
    try:
        riecken_records.build_records(package, data, "1590", CLEARING_NAME, {}, "Eingangsrechnungen", "Rechnungen Nachlauf")
    except ValueError as exc:
        assert "weicht von account_config.clarification" in str(exc)
    else:
        raise AssertionError("abweichendes Klärungskonto wurde akzeptiert")

    overrides = {"V0004": {"topic": "Anlage GWG", "target": "Ziel 0480 Dyson AM07"}}
    work = root / "work"
    result = subprocess.run(
        [sys.executable, str(SCRIPT_DIR / "riecken_records.py"), "--package", str(package), "--input", str(scenario.root / "riecken.json"),
         "--clearing-account", CLEARING, "--clearing-name", CLEARING_NAME, "--output", str(work), "--klaer-texte", str(write_json(root / "klaer.json", overrides))],
        capture_output=True, text=True, encoding="utf-8",
    )
    assert result.returncode == 0, result.stderr
    records = json.loads((work / "riecken_records.json").read_text(encoding="utf-8"))
    rejected = json.loads((work / "nicht_uebertragbar.json").read_text(encoding="utf-8"))

    green = records["records_gruen"]
    red = records["records_rot"]
    assert len(green) == 1 and len(red) == 1
    green_records = [item for call in green[0]["calls"] for item in call]
    red_records = [item for call in red[0]["calls"] for item in call]
    assert len(green_records) == 2 and {item["transaction_id"] for item in green_records} == {"V0001", "V0002"}
    assert green[0]["description"] == "Eingangsrechnungen" and red[0]["description"] == "Rechnungen Nachlauf"
    assert "Klärungsposten" not in json.dumps(records["preview"], ensure_ascii=False)
    # Ein Riecken-Stapelname mit Klärungs- oder Prüfhinweis wird abgewiesen.
    try:
        riecken_records.build_records(package, data, CLEARING, CLEARING_NAME, {}, "Eingangsrechnungen", "Klärungsposten")
    except ValueError as exc:
        assert "kein Hinweis auf Klärung" in str(exc)
    else:
        raise AssertionError("Riecken-Stapelname Klärungsposten wurde akzeptiert")
    # grün 1:1, Gutschrift mit vertauschtem Soll/Haben
    by_tid = {item["transaction_id"]: item for item in green_records}
    assert by_tid["V0001"]["debit_account"] == "4900" and by_tid["V0001"]["credit_account"] == "70001" and by_tid["V0001"]["tax_key"] == 401
    assert by_tid["V0002"]["debit_account"] == "70001" and by_tid["V0002"]["credit_account"] == "4900" and by_tid["V0002"]["amount"] == 50.0
    assert all(item["date"] == "2025-12-15" and item.get("document_guid") for item in green_records)
    assert Decimal(green[0]["total_amount"]) == extf_total(datev / "EXTF_Buchungsstapel_2025-12.csv")

    # rot: Klärungskonto, Datum, GUID, KLÄR-Text
    assert {item["transaction_id"] for item in red_records} == {"V0003", "V0004"}
    for item in red_records:
        assert item["debit_account"] == CLEARING and item["credit_account"] == "70001", item
        assert item["clearing_side"] == "debit"
        assert item["date"] == "2025-12-15" and item.get("document_guid")
        assert item["posting_description"].startswith(f"KLÄR {item['transaction_id']} ")
        assert len(item["posting_description"]) <= 60
        assert not (item["debit_account"] == CLEARING and item["credit_account"] == CLEARING)
        assert Decimal(str(item["amount"])) != Decimal("0.01") and item["amount"] > 0
    red_by_tid = {item["transaction_id"]: item for item in red_records}
    assert red_by_tid["V0004"]["posting_description"] == "KLÄR V0004 Anlage GWG Ziel 0480 Dyson AM07"
    assert "Konto offen" in red_by_tid["V0003"]["posting_description"]
    # Summenabgleich: übertragbar + nicht übertragbar = EXTF
    red_extf = extf_total(datev / "EXTF_Klaerungsposten_2025-12.csv")
    assert Decimal(red[0]["total_amount"]) + Decimal(red[0]["not_transferred_total"]) == red_extf
    assert red[0]["extf_rows"] == 5 and red[0]["record_count"] == 2 and red[0]["not_transferred_count"] == 3
    reasons = {item["transaction_id"]: item["reason"] for item in rejected}
    assert reasons == {"V0005": riecken_records.REASON_BOTH_SIDES, "V0006": riecken_records.REASON_AMOUNT, "V0007": riecken_records.REASON_DATE}, reasons
    assert all(item["decision"] is None and item["decided_by"] is None for item in rejected)
    # Beleg-GUIDs stammen aus document.xml
    guids = riecken_records.document_guids(package)
    assert all(item["document_guid"] in guids for item in green_records + red_records)
    return package, work, records, rejected


def write_json(path: Path, payload) -> Path:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


# 3. Excel, Statusdatei, Verschieben -------------------------------------------------

def test_workbook_update(package: Path, work: Path, records: dict, rejected: list) -> None:
    transfer = {
        "transferred_at": "2026-10-08T15:30:00+02:00",
        "clearing_account": {"account": CLEARING, "name": CLEARING_NAME},
        "change_plans": [
            {"id": "aaaa1111", "type": "posting_batch", "executed_at": "2026-10-08T15:20:00+02:00"},
            {"id": "bbbb2222", "type": "posting_batch", "executed_at": "2026-10-08T15:28:00+02:00"},
        ],
        "sequences": [
            {**records["preview"][0], "change_plan_id": "aaaa1111", "source_file": records["records_gruen"][0]["source_file"]},
            {**records["preview"][1], "change_plan_id": "bbbb2222", "source_file": records["records_rot"][0]["source_file"]},
        ],
        "not_transferred": [
            {"transaction_id": "V0005", "reason": "beide Kontoseiten offen", "decision": "nicht buchen, Beleg bleibt im Belegtransfer", "decided_by": "Oliver Burchardt", "decided_at": "2026-10-08"},
            {"transaction_id": "V0006", "reason": "Betrag unsicher", "decision": "nicht buchen, vollständige Rechnung anfordern", "decided_by": "Oliver Burchardt", "decided_at": "2026-10-08"},
        ],
        "removed_documents": [],
        "transferred_files": records["transferred_files_expected"],
        "visibility_check": {"tool": "datev_get_account_postings", "result": "nicht sichtbar", "checked_at": "2026-10-08T15:35:00+02:00"},
        "open_steps": ["Belegtransfer-ZIPs in DUO hochladen"],
    }
    transfer_path = write_json(work / "riecken_transfer.json", transfer)
    result = subprocess.run(
        [sys.executable, str(SCRIPT_DIR / "update_review_workbook_riecken.py"), "--package", str(package),
         "--records", str(work / "riecken_records.json"), "--transfer", str(transfer_path)],
        capture_output=True, text=True, encoding="utf-8",
    )
    assert result.returncode == 0, result.stderr
    replaced = package / "03_Technische_Protokolle" / "ersetzt"
    workbook_path = next(path for path in (package / "02_Buchungspruefung").glob("Buchungspruefung_*.xlsx"))
    assert (replaced / f"{workbook_path.stem}_vor_Riecken.xlsx").is_file()
    for name in records["transferred_files_expected"]:
        assert not (package / "01_DATEV_Import" / name).exists() and (replaced / name).is_file(), name
    assert list((package / "01_DATEV_Import").glob("Belegtransfer_*.zip")), "Belegtransfer-ZIPs bleiben im Importordner"

    wb = load_workbook(workbook_path)
    assert wb.sheetnames.index("Riecken-Übertragung") == wb.sheetnames.index("Übersicht") + 1
    review = wb["Belegprüfung"]
    cols = {str(c.value): c.column for c in review[1]}
    rows = {str(review.cell(r, cols["Vorgangs-ID"]).value): r for r in range(2, review.max_row + 1)}
    assert review.cell(rows["V0001"], cols["Buchungsstapel"]).value == "Riecken: Eingangsrechnungen 2025-12"
    assert review.cell(rows["V0003"], cols["Buchungsstapel"]).value == f"Riecken: Rechnungen Nachlauf 2025-12 (Konto {CLEARING})"
    assert str(review.cell(rows["V0003"], cols["Kontierung"]).value).startswith("Gebucht über Riecken: Soll 1599 an Haben 70001")
    assert "Kontierung aus Beleg nicht erkennbar" in str(review.cell(rows["V0003"], cols["Offener Punkt / nächster Schritt"]).value)
    assert str(review.cell(rows["V0003"], cols["Offener Punkt / nächster Schritt"]).value).startswith("Umbuchung: Klärungskonto 1599")
    assert str(review.cell(rows["V0005"], cols["Buchungsstapel"]).value).startswith("nicht gebucht (Entscheidung Oliver Burchardt")
    assert review.cell(rows["V0005"], cols["Bearbeitungsstatus"]).value == "nicht übernommen"
    assert "Oliver Burchardt, 2026-10-08" in str(review.cell(rows["V0005"], cols["Mitarbeiter-Ergebnis"]).value)
    assert str(review.cell(rows["V0007"], cols["Buchungsstapel"]).value).startswith("nicht gebucht (")
    assert review.cell(rows["V0007"], cols["Bearbeitungsstatus"]).value == "nicht übernommen"

    bookings = wb["Buchungszeilen"]
    bcols = {str(c.value): c.column for c in bookings[1]}
    brows = {str(bookings.cell(r, bcols["Vorgangs-ID"]).value): r for r in range(2, bookings.max_row + 1)}
    assert bookings.cell(brows["V0003"], bcols["Konto"]).value == CLEARING
    assert bookings.cell(brows["V0003"], bcols["Kontobezeichnung"]).value == CLEARING_NAME
    assert str(bookings.cell(brows["V0003"], bcols["Buchungstext"]).value).startswith("KLÄR V0003 ")
    assert bookings.cell(brows["V0004"], bcols["Buchungstext"]).value == "KLÄR V0004 Anlage GWG Ziel 0480 Dyson AM07"
    assert bookings.cell(brows["V0002"], bcols["Buchungsstapel"]).value == "Riecken: Eingangsrechnungen 2025-12"
    assert bookings.cell(brows["V0002"], bcols["Konto"]).value == "4900"

    summary = wb["Übersicht"]
    texts = [str(summary.cell(r, 1).value) for r in range(1, summary.max_row + 1)]
    assert "Riecken-Stapel" in texts and "Rechnungen Nachlauf" in texts and "DATEV-Stapel" not in texts
    guide = wb["Anleitung"]
    guide_texts = " ".join(str(guide.cell(r, c).value or "") for r in range(1, guide.max_row + 1) for c in (1, 2))
    assert "Zielsaldo 0" in guide_texts and "Klärungskonto" in guide_texts and "dürfen nicht zusätzlich importiert werden" in guide_texts
    sheet = wb["Riecken-Übertragung"]
    sheet_text = " ".join(str(sheet.cell(r, c).value or "") for r in range(1, sheet.max_row + 1) for c in range(1, 7))
    assert "aaaa1111" in sheet_text and "bbbb2222" in sheet_text and "V0005" in sheet_text and "V0007" in sheet_text and "Belegtransfer-ZIPs in DUO hochladen" in sheet_text

    status = (package / "00_STATUS_NACH_RIECKEN_UEBERTRAGUNG.md").read_text(encoding="utf-8")
    for needle in ("aaaa1111", "bbbb2222", "Rechnungen Nachlauf", "V0005", "V0006", "V0007", "nicht sichtbar", "EXTF_Klaerungsposten_2025-12.csv", "Offene Schritte"):
        assert needle in status, needle
    manifest = json.loads((package / "03_Technische_Protokolle" / "Laufmanifest.json").read_text(encoding="utf-8"))
    assert manifest["riecken_transfer"]["change_plans"][0]["id"] == "aaaa1111"


# 4. Beleg entfernen --------------------------------------------------------------

def test_remove_document(package: Path) -> None:
    index = json.loads((package / "03_Technische_Protokolle" / "Belegindex.json").read_text(encoding="utf-8"))
    entry = next(item for item in index if item.get("transaction_ids") == ["V0005"] or item.get("transaction_id") == "V0005")
    archive = package / "01_DATEV_Import" / entry["document_package"]
    guid = entry["document_guid"]
    with zipfile.ZipFile(archive) as bundle:
        before = set(bundle.namelist())
    log = remove_document_from_transfer.remove_document(archive, guid, reason="Fehlbuchung des Mandanten bereits korrigiert")
    assert log["testzip"] == "ok" and log["removed_file"] == entry["technical_filename"]
    assert Path(log["original_moved_to"]).is_file() and Path(log["original_moved_to"]).parent.name == "ersetzt"
    with zipfile.ZipFile(archive) as bundle:
        after = set(bundle.namelist())
        assert bundle.testzip() is None
        xml = ET.fromstring(bundle.read("document.xml"))
    assert before - after == {entry["technical_filename"]}
    assert all(str(doc.get("guid", "")).upper() != guid for doc in xml.iter() if doc.tag.endswith("}document"))
    assert sum(1 for doc in xml.iter() if doc.tag.endswith("}document")) == len(after) - 1
    try:
        remove_document_from_transfer.remove_document(archive, guid)
    except ValueError as exc:
        assert "nicht in document.xml" in str(exc)
    else:
        raise AssertionError("zweite Entfernung derselben GUID wurde akzeptiert")


def main() -> None:
    test_version_and_texts()
    with tempfile.TemporaryDirectory(prefix="bk_v150_") as name:
        root = Path(name)
        package, work, records, rejected = test_records(root)
        test_workbook_update(package, work, records, rejected)
        test_remove_document(package)
    print("v1.5.0 Riecken-Übertragung, Klärungskonto, KLÄR-Texte, Excel-Stand, Statusdatei, Belegentfernung und Version: OK")


if __name__ == "__main__":
    main()
