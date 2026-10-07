from __future__ import annotations

import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from evaluate_review_return import REQUIRED_HEADERS, evaluate


def build_workbook(path: Path, rows: list[tuple[str, str, str]]) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Übersicht"
    summary["A1"] = "Übersicht 13402 – 2026-07"
    review = wb.create_sheet("Belegprüfung")
    review.append(REQUIRED_HEADERS)
    for position, (transaction_id, light, partner) in enumerate(rows, start=1):
        review.append(
            [
                light,
                "EXTF_Buchungsstapel_2026-07.csv",
                transaction_id,
                datetime(2026, 7, min(position, 28)),
                partner,
                f"BF-{position}",
                100.0 + position,
                "EUR",
                "2026-07",
                "1000 an 70000",
                "Belegbezogene Ableitung",
                "Fachliche Prüfung erforderlich",
                "Behandlung abschließend dokumentieren",
                "offen",
                "",
            ]
        )
    bookings = wb.create_sheet("Buchungszeilen")
    bookings.append(["Vorgangs-ID", "Kontierung"])
    for transaction_id, _, _ in rows:
        bookings.append([transaction_id, "1000 an 70000"])
    wb.save(path)


def return_copy(
    original: Path,
    returned: Path,
    values: dict[str, tuple[str, str]],
) -> None:
    shutil.copyfile(original, returned)
    wb = load_workbook(returned)
    ws = wb["Belegprüfung"]
    for row in range(2, ws.max_row + 1):
        transaction_id = str(ws.cell(row, 3).value)
        if transaction_id in values:
            ws.cell(row, 14).value, ws.cell(row, 15).value = values[transaction_id]
    wb.save(returned)


def run_case(
    root: Path,
    name: str,
    rows: list[tuple[str, str, str]],
    values: dict[str, tuple[str, str]],
    profile: Path | None = None,
    register: Path | None = None,
) -> tuple[dict, Path, Path]:
    original = root / f"{name}-original.xlsx"
    returned = root / f"{name}-returned.xlsx"
    output = root / f"{name}-output"
    build_workbook(original, rows)
    return_copy(original, returned, values)
    return (
        evaluate(original, returned, output, profile, register),
        original,
        returned,
    )


def main() -> None:
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)

        result, _, _ = run_case(
            root,
            "complete",
            [("V0001", "Rot", "A GmbH"), ("V0002", "Rot", "B GmbH")],
            {
                "V0001": ("geändert", "Als Aufwand gebucht"),
                "V0002": ("unverändert übernommen", ""),
            },
        )
        assert result["status"] == "vollständig"
        assert Path(result["output"]).is_file()

        result, _, _ = run_case(
            root,
            "open",
            [("V0001", "Rot", "A GmbH"), ("V0002", "Rot", "B GmbH")],
            {"V0001": ("unverändert übernommen", "")},
        )
        assert result["status"] == "unvollständig"
        assert any("V0002" in error and "offen" in error for error in result["errors"])

        result, _, _ = run_case(
            root,
            "missing-result",
            [("V0001", "Rot", "A GmbH")],
            {"V0001": ("geändert", "")},
        )
        assert result["status"] == "unvollständig"
        assert any("Mitarbeiter-Ergebnis" in error for error in result["errors"])

        result, _, returned = run_case(
            root,
            "integrity",
            [("V0001", "Rot", "A GmbH")],
            {"V0001": ("unverändert übernommen", "")},
        )
        wb = load_workbook(returned)
        wb["Belegprüfung"].cell(2, 10).value = "1200 an 70000"
        wb.save(returned)
        result = evaluate(root / "integrity-original.xlsx", returned, root / "integrity-output-2")
        assert result["status"] == "unvollständig"
        assert not result["integrity"]
        assert any("Kontierung" in error for error in result["errors"])

        result, _, _ = run_case(
            root,
            "profile",
            [("V0001", "Rot", "A GmbH")],
            {"V0001": ("geändert", "Künftig immer auf Spezialkonto buchen")},
        )
        assert len(result["profile_suggestions"]) == 1

        result, original, returned = run_case(
            root,
            "accrual",
            [("V0001", "Rot", "AXA Versicherung")],
            {"V0001": ("geändert", "Auf ARAP erfasst und abzugrenzen")},
        )
        assert len(result["register_suggestions"]) == 1
        register = root / "register.md"
        register.write_text("AXA Versicherung – ARAP 2026", encoding="utf-8")
        result = evaluate(original, returned, root / "accrual-output-2", None, register)
        assert not result["register_suggestions"]

        result, _, _ = run_case(
            root,
            "one-time",
            [("V0001", "Rot", "BOBE Tiefbau")],
            {"V0001": ("geändert", "Geprüft, als Instandhaltung gebucht")},
        )
        assert not result["profile_suggestions"]
        assert not result["register_suggestions"]

        product_rows = [
            ("V0023", "Rot", "BOBE Tiefbau"),
            ("V0060", "Rot", "AXA Versicherung"),
            *[(f"V01{number:02d}", "Rot", f"Lieferant {number}") for number in range(1, 10)],
        ]
        product_values = {
            "V0023": ("geändert", "Geprüft, als Instandhaltung gebucht"),
            "V0060": (
                "geändert",
                "Auf ARAP erfasst, Abgrenzung auf 980 im Juli gebucht, ab 9/26 über 12 Monate abzugrenzen",
            ),
            **{
                f"V01{number:02d}": ("unverändert übernommen", "")
                for number in range(1, 10)
            },
        }
        profile = root / "profile.md"
        profile.write_text("BOBE Tiefbau; AXA Versicherung; Personenkonten geprüft", encoding="utf-8")
        result, _, _ = run_case(
            root,
            "product-13402",
            product_rows,
            product_values,
            profile,
            register,
        )
        assert result["status"] == "vollständig"
        assert result["counts"]["Rot"] == {"gesamt": 11, "abgeschlossen": 11}
        assert not result["profile_suggestions"]
        assert not result["register_suggestions"]

    print("review return tests: OK")


if __name__ == "__main__":
    main()
