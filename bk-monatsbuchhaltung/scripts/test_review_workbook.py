from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from build_review_workbook import TRAFFIC_COLORS, build


def document(light: str, number: int) -> dict:
    return {
        "transaction_id": f"V{number:04d}",
        "recognized_date": "2026-07-15",
        "partner": f"Testpartner {number}",
        "invoice_number": f"RE-{number}",
        "total_amount": "119.00",
        "currency": "EUR",
        "period": "2026-07",
        "processing_status": "Buchungszeile erzeugt",
        "traffic_light": light,
        "derivation": "Testkontierung",
        "reason": "Testbegründung",
        "bookings": [
            {
                "amount": "119.00",
                "debit_credit": "S",
                "account": "4900",
                "account_name": "Sonstiger Aufwand",
                "contra_account": "70001",
                "contra_account_name": f"Testpartner {number}",
                "bu_key": "401",
                "document_field_1": f"RE-{number}",
                "booking_text": "Testbuchung",
            }
        ],
    }


def assert_traffic_fills(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    for sheet_name in ("Belegprüfung", "Buchungszeilen"):
        sheet = workbook[sheet_name]
        observed = {}
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, 1)
            if cell.value in TRAFFIC_COLORS:
                observed[cell.value] = cell.fill.fgColor.rgb[-6:]
                if cell.fill.fill_type != "solid":
                    raise AssertionError(
                        f"{sheet_name}!A{row}: Ampelfarbe ist nicht als "
                        "solide Zellfüllung hinterlegt"
                    )
                if cell.font.bold is not True:
                    raise AssertionError(
                        f"{sheet_name}!A{row}: Ampelstatus ist nicht fett"
                    )
        if observed != TRAFFIC_COLORS:
            raise AssertionError(
                f"{sheet_name}: Ampelfarben {observed!r} statt "
                f"{TRAFFIC_COLORS!r}"
            )


def main() -> None:
    data = {
        "run": {"mandantennummer": "12345", "buchungsmonat": "2026-07"},
        "documents": [
            document("Grün", 1),
            document("Rot", 3),
        ],
    }
    with tempfile.TemporaryDirectory(prefix="bk_review_workbook_") as temp_name:
        output = Path(temp_name) / "Buchungspruefung.xlsx"
        build(data, output)
        assert_traffic_fills(output)
    print("review workbook tests: OK")


if __name__ == "__main__":
    main()
