"""Prüfungs-Excel und Paket nach der Riecken-Übertragung auf den Riecken-Stand bringen.

Umsetzung von SKILL.md Abschnitt 5, Absatz „Nach der Übertragung“:

- Ursprungsfassung der Excel als ``<Name>_vor_Riecken.xlsx`` unter ``03_Technische_Protokolle/ersetzt/``,
- Spalte ``Buchungsstapel`` in ``Belegprüfung`` und ``Buchungszeilen`` auf die Riecken-Stapel,
- rote Zeilen mit Klärungskonto und tatsächlich gebuchtem KLÄR-Text,
- Nutzerentscheidungen zu nicht übertragenen Vorgängen (``nicht übernommen`` + Entscheider, Datum, Begründung),
- Übersicht mit Riecken-Stapeln, Anleitung mit Klärungskonto, neues Blatt ``Riecken-Übertragung``,
- Statusdatei ``00_STATUS_NACH_RIECKEN_UEBERTRAGUNG.md`` im Paketstamm,
- Verschieben der übertragenen EXTF-Dateien nach ``03_Technische_Protokolle/ersetzt/``,
- Übernahme von ``riecken_transfer`` in das Laufmanifest.

Eingaben: Paketordner, ``riecken_records.json`` (aus ``riecken_records.py``) und
``riecken_transfer.json`` (Schema: references/EINGABESCHEMA.md, Abschnitt ``riecken_transfer``).
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

DATEV_FOLDER = "01_DATEV_Import"
REVIEW_FOLDER = "02_Buchungspruefung"
LOG_FOLDER = "03_Technische_Protokolle"
REPLACED_FOLDER = "ersetzt"
STATUS_FILE = "00_STATUS_NACH_RIECKEN_UEBERTRAGUNG.md"
TRANSFER_SHEET = "Riecken-Übertragung"
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
NOT_BOOKED = "nicht gebucht"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--package", required=True, type=Path, help="Paketordner (enthält 01_DATEV_Import, 02_Buchungspruefung, 03_Technische_Protokolle)")
    parser.add_argument("--records", required=True, type=Path, help="riecken_records.json aus riecken_records.py")
    parser.add_argument("--transfer", required=True, type=Path, help="riecken_transfer.json mit Change-Plan-IDs, Stapeln und Entscheidungen")
    parser.add_argument("--workbook", type=Path, help="Prüfungs-Excel; Standard: einzige Buchungspruefung_*.xlsx in 02_Buchungspruefung")
    parser.add_argument("--keep-extf", action="store_true", help="Übertragene EXTF-Dateien nicht nach ersetzt/ verschieben")
    return parser.parse_args()


def money(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


def fmt_money(value: Any) -> str:
    amount = money(value)
    text = f"{amount:,.2f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def column_map(sheet) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in sheet[1] if cell.value}


def find_workbook(package: Path, explicit: Path | None) -> Path:
    if explicit:
        return explicit
    candidates = sorted((package / REVIEW_FOLDER).glob("Buchungspruefung_*.xlsx"))
    candidates = [path for path in candidates if "_vor_Riecken" not in path.name]
    if len(candidates) != 1:
        raise ValueError(f"Prüfungs-Excel nicht eindeutig in {package / REVIEW_FOLDER}: {[c.name for c in candidates]}")
    return candidates[0]


def record_index(records: dict[str, Any]) -> tuple[dict[tuple[str, int], dict[str, Any]], dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    """(file, csv_row) -> record; transaction_id -> sequence; transaction_id -> first red record."""
    by_row: dict[tuple[str, int], dict[str, Any]] = {}
    by_transaction: dict[str, dict[str, Any]] = {}
    red_by_transaction: dict[str, dict[str, Any]] = {}
    for kind in ("records_gruen", "records_rot"):
        for sequence in records.get(kind, []):
            for call in sequence.get("calls", []):
                for record in call:
                    key = (str(record.get("source_file")), int(record.get("csv_row", 0)))
                    by_row[key] = record
                    tid = str(record.get("transaction_id"))
                    by_transaction.setdefault(tid, sequence)
                    if kind == "records_rot":
                        red_by_transaction.setdefault(tid, record)
    return by_row, by_transaction, red_by_transaction


def batch_text(sequence: dict[str, Any], clearing: dict[str, Any]) -> str:
    if sequence.get("kind") == "rot":
        return f"Riecken: {sequence['description']} {sequence['month']} (Konto {clearing.get('account')})"
    return f"Riecken: {sequence['description']} {sequence['month']}"


def not_booked_text(entry: dict[str, Any] | None, reason: str = "") -> str:
    if entry:
        who = entry.get("decided_by") or "Auftraggeber"
        return f"{NOT_BOOKED} (Entscheidung {who}: {entry.get('decision') or entry.get('reason') or reason})"
    return f"{NOT_BOOKED} ({reason or 'nicht übertragbar'})"


def update_bookings_sheet(sheet, by_row, by_transaction, red_by_transaction, decisions, rejected_by_tid, clearing) -> None:
    cols = column_map(sheet)
    stapel = cols["Buchungsstapel"]
    tid_col = cols["Vorgangs-ID"]
    for row in range(2, sheet.max_row + 1):
        tid = str(sheet.cell(row, tid_col).value or "")
        file_name = str(sheet.cell(row, stapel).value or "")
        record = None
        for key, candidate in by_row.items():
            if key[0] == file_name and str(candidate.get("transaction_id")) == tid and candidate is not None and not candidate.get("_used"):
                record = candidate
                candidate["_used"] = True
                break
        if record is None:
            if tid in decisions or tid in rejected_by_tid:
                sheet.cell(row, stapel).value = not_booked_text(decisions.get(tid), (rejected_by_tid.get(tid) or {}).get("reason", ""))
            continue
        sequence = by_transaction[tid]
        sheet.cell(row, stapel).value = batch_text(sequence, clearing)
        if sequence.get("kind") != "rot":
            continue
        side = record.get("clearing_side")
        if side == "debit":
            target_cols = ("Konto", "Kontobezeichnung") if str(sheet.cell(row, cols["Soll/Haben"]).value or "S").upper() == "S" else ("Gegenkonto", "Gegenkontobezeichnung")
        elif side == "credit":
            target_cols = ("Gegenkonto", "Gegenkontobezeichnung") if str(sheet.cell(row, cols["Soll/Haben"]).value or "S").upper() == "S" else ("Konto", "Kontobezeichnung")
        else:
            target_cols = None
        if target_cols:
            sheet.cell(row, cols[target_cols[0]]).value = clearing.get("account")
            sheet.cell(row, cols[target_cols[1]]).value = clearing.get("name")
        sheet.cell(row, cols["Buchungstext"]).value = record.get("posting_description")
    for candidate in by_row.values():
        candidate.pop("_used", None)


def update_review_sheet(sheet, by_transaction, red_by_transaction, decisions, rejected_by_tid, clearing) -> None:
    cols = column_map(sheet)
    stapel = cols["Buchungsstapel"]
    tid_col = cols["Vorgangs-ID"]
    for row in range(2, sheet.max_row + 1):
        tid = str(sheet.cell(row, tid_col).value or "")
        if not tid:
            continue
        sequence = by_transaction.get(tid)
        if sequence is None:
            if tid in decisions or tid in rejected_by_tid:
                entry = decisions.get(tid)
                sheet.cell(row, stapel).value = not_booked_text(entry, (rejected_by_tid.get(tid) or {}).get("reason", ""))
                sheet.cell(row, cols["Bearbeitungsstatus"]).value = "nicht übernommen"
                if entry:
                    sheet.cell(row, cols["Mitarbeiter-Ergebnis"]).value = (
                        f"{entry.get('decided_by') or 'Auftraggeber'}, {entry.get('decided_at') or ''}: "
                        f"{entry.get('decision') or ''} – {entry.get('reason') or ''}"
                    ).strip(" –:")
            continue
        sheet.cell(row, stapel).value = batch_text(sequence, clearing)
        if sequence.get("kind") != "rot":
            continue
        record = red_by_transaction.get(tid)
        if not record:
            continue
        actual = (
            f"Gebucht über Riecken: Soll {record.get('debit_account')} an Haben {record.get('credit_account')}, "
            f"{fmt_money(record.get('amount'))} EUR, Text „{record.get('posting_description')}“."
        )
        instruction = (
            f"Umbuchung: Klärungskonto {clearing.get('account')} {clearing.get('name')} nach Klärung auf das Zielkonto umbuchen; "
            "Klärungskonto für die Periode auf Saldo 0 bringen, erst danach Klärungsstapel festschreiben."
        )
        kont = sheet.cell(row, cols["Kontierung"])
        kont.value = f"{actual}\n{kont.value or ''}".rstrip()
        nxt = sheet.cell(row, cols["Offener Punkt / nächster Schritt"])
        nxt.value = f"{instruction}\n{nxt.value or ''}".rstrip()


def update_summary_sheet(sheet, transfer: dict[str, Any], records: dict[str, Any], decisions: dict[str, Any]) -> None:
    header_row = None
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "DATEV-Stapel":
            header_row = row
            break
    if header_row is None:
        header_row = sheet.max_row + 2
    sheet.delete_rows(header_row, sheet.max_row - header_row + 1)
    sheet.cell(header_row, 1).value = "Riecken-Stapel"
    sheet.cell(header_row, 2).value = "Monat"
    sheet.cell(header_row, 3).value = "Zeilen"
    sheet.cell(header_row, 4).value = "Summe"
    for col in range(1, 5):
        sheet.cell(header_row, col).font = Font(bold=True)
        sheet.cell(header_row, col).fill = HEADER_FILL
    row = header_row + 1
    for sequence in transfer.get("sequences", []):
        sheet.cell(row, 1).value = sequence.get("description")
        sheet.cell(row, 2).value = sequence.get("month")
        sheet.cell(row, 3).value = int(sequence.get("record_count", 0))
        sheet.cell(row, 4).value = float(money(sequence.get("total_amount")))
        sheet.cell(row, 3).number_format = "#,##0"
        sheet.cell(row, 4).number_format = "#,##0.00"
        row += 1
    for r in range(1, header_row):
        if sheet.cell(r, 1).value == "Offene Klärfälle":
            open_cases = sum(len(seq.get("calls", [])) and sum(len(call) for call in seq.get("calls", [])) for seq in records.get("records_rot", []))
            sheet.cell(r, 2).value = open_cases
            sheet.cell(r, 1).value = "Offene Klärfälle (Riecken-Stapel Klärungsposten, ohne entschiedene Fälle)"
        if sheet.cell(r, 1).value == "Kontrollpunkt":
            pass
    sheet.cell(row + 1, 1).value = "Nicht gebuchte Vorgänge (Entscheidung Auftraggeber)"
    sheet.cell(row + 1, 2).value = len(decisions)


def update_guide_sheet(sheet, clearing: dict[str, Any]) -> None:
    for row in range(2, sheet.max_row + 1):
        key = str(sheet.cell(row, 1).value or "")
        if key == "1. Belegprüfung":
            sheet.cell(row, 2).value = (
                "Rote Fälle bearbeiten. Grüne Vorgänge wurden über den Riecken-Connector im Stapel Eingangsrechnungen je Monat gebucht, "
                f"rote im Stapel Klärungsposten je Monat auf das Klärungskonto {clearing.get('account')} {clearing.get('name')}. "
                "Rechts neben der Ampel steht der Riecken-Stapel. Nicht gebuchte Vorgänge tragen den Grund der Entscheidung."
            )
        elif key == "2. Buchungszeilen":
            sheet.cell(row, 2).value = (
                "Direkt rechts neben der Ampel steht der Riecken-Stapel; bei roten Zeilen stehen Klärungskonto und der tatsächlich gebuchte "
                "KLÄR-Buchungstext (KLÄR <Vorgangs-ID> <offenes Thema> <Zielkonto>)."
            )
        elif key == "DATEV-Import":
            sheet.cell(row, 1).value = "DATEV / DUO"
            sheet.cell(row, 2).value = (
                "Buchungsstapel, Klärungsposten und Personenkonten wurden über den Riecken-Connector nach DATEV geschrieben; die EXTF-Dateien "
                "dürfen nicht zusätzlich importiert werden (sie liegen unter 03_Technische_Protokolle/ersetzt). Reguläre und Avis-Belegtransfer-ZIPs "
                "in DATEV Unternehmen online hochladen; erst danach lösen die Beleglinks der Riecken-Buchungen auf."
            )
    sheet.append([
        "Klärungskonto",
        f"{clearing.get('account')} {clearing.get('name')} ist ein reines Arbeitskonto mit Zielsaldo 0. Jede rote Zeile vom Klärungskonto auf das "
        "Zielkonto umbuchen (Anlagen zuerst in der Anlagenbuchführung anlegen); erst wenn das Konto für die Periode den Saldo 0 hat, "
        "den Klärungsstapel festschreiben.",
    ])
    last = sheet.max_row
    sheet.cell(last, 1).font = Font(bold=True)
    sheet.cell(last, 2).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[last].height = 48


def build_transfer_sheet(workbook, transfer: dict[str, Any], rejected: list[dict[str, Any]]) -> None:
    if TRANSFER_SHEET in workbook.sheetnames:
        del workbook[TRANSFER_SHEET]
    position = workbook.sheetnames.index("Übersicht") + 1 if "Übersicht" in workbook.sheetnames else 1
    sheet = workbook.create_sheet(TRANSFER_SHEET, position)
    clearing = transfer.get("clearing_account") or {}
    sheet.append(["Schritt", "Inhalt", "Zeilen", "Summe", "Change-Plan-ID", "Status"])
    plans = {str(plan.get("id")): plan for plan in transfer.get("change_plans", [])}
    partner_plans = [plan for plan in transfer.get("change_plans", []) if plan.get("type") == "business_partner"]
    for plan in partner_plans:
        sheet.append(["Personenkonten", "Neue Debitoren/Kreditoren aus EXTF_Debitoren_Kreditoren.csv", "", "", plan.get("id"), "ausgeführt" if plan.get("executed_at") else "offen"])
    for sequence in transfer.get("sequences", []):
        plan = plans.get(str(sequence.get("change_plan_id")), {})
        sheet.append([
            f"{sequence.get('description')} {sequence.get('month')}",
            sequence.get("source_file") or "",
            int(sequence.get("record_count", 0)),
            float(money(sequence.get("total_amount"))),
            sequence.get("change_plan_id") or "",
            "ausgeführt" if plan.get("executed_at") or sequence.get("status") == "ausgeführt" else (sequence.get("status") or "offen"),
        ])
    sheet.append([])
    sheet.append(["Klärungskonto", f"{clearing.get('account')} {clearing.get('name')} (Arbeitskonto, Zielsaldo 0)", "", "", "", ""])
    sheet.append([])
    sheet.append(["Nicht gebucht", "Grund", "Betrag", "Entscheidung", "Entscheider", "Datum"])
    decided = {str(item.get("transaction_id")): item for item in transfer.get("not_transferred", [])}
    for item in rejected:
        entry = decided.get(str(item.get("transaction_id")), {})
        sheet.append([
            item.get("transaction_id"), item.get("reason"), float(money(item.get("amount"))) if item.get("amount") is not None else "",
            entry.get("decision") or "offen", entry.get("decided_by") or "", entry.get("decided_at") or "",
        ])
    for tid, entry in decided.items():
        if not any(str(item.get("transaction_id")) == tid for item in rejected):
            sheet.append([tid, entry.get("reason"), "", entry.get("decision") or "", entry.get("decided_by") or "", entry.get("decided_at") or ""])
    if transfer.get("removed_documents"):
        sheet.append([])
        sheet.append(["Entfernte Belege", "GUID", "Belegtransfer", "Grund", "", ""])
        for item in transfer["removed_documents"]:
            sheet.append(["", item.get("guid"), item.get("file"), item.get("reason"), "", ""])
    sheet.append([])
    sheet.append(["Offene Schritte", "", "", "", "", ""])
    for step in transfer.get("open_steps", []):
        sheet.append(["", step, "", "", "", ""])
    visibility = transfer.get("visibility_check") or {}
    if visibility:
        sheet.append(["Sichtbarkeit", f"{visibility.get('tool')}: {visibility.get('result')} ({visibility.get('checked_at')})", "", "", "", ""])
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, 1).value in {"Schritt", "Nicht gebucht", "Entfernte Belege", "Offene Schritte"}:
            for col in range(1, 7):
                sheet.cell(row, col).font = Font(bold=True)
                sheet.cell(row, col).fill = HEADER_FILL
        if isinstance(sheet.cell(row, 3).value, int):
            sheet.cell(row, 3).number_format = "#,##0"
        if isinstance(sheet.cell(row, 4).value, float):
            sheet.cell(row, 4).number_format = "#,##0.00"
        if isinstance(sheet.cell(row, 3).value, float):
            sheet.cell(row, 3).number_format = "#,##0.00"
        for col in (2, 4):
            sheet.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
    for letter, width in zip("ABCDEF", (34, 60, 12, 16, 36, 16)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"


def status_markdown(package: Path, transfer: dict[str, Any], records: dict[str, Any], rejected: list[dict[str, Any]], moved: list[str]) -> str:
    clearing = transfer.get("clearing_account") or records.get("clearing_account") or {}
    lines = [
        f"# Status nach Riecken-Übertragung – Mandant {records.get('mandant')}",
        "",
        f"Stand: {transfer.get('transferred_at') or datetime.now().replace(microsecond=0).isoformat()}",
        f"Klärungskonto: {clearing.get('account')} {clearing.get('name')} (Arbeitskonto, Zielsaldo 0)",
        "",
        "## Change-Pläne",
        "",
        "| ID | Typ | ausgeführt am |",
        "|---|---|---|",
    ]
    for plan in transfer.get("change_plans", []):
        lines.append(f"| {plan.get('id')} | {plan.get('type')} | {plan.get('executed_at') or 'offen'} |")
    lines += ["", "## Stapel je Monat", "", "| Monat | Stapel | Zeilen | Summe EUR | Change-Plan | EXTF-Quelle |", "|---|---|---|---|---|---|"]
    for sequence in transfer.get("sequences", []):
        lines.append(
            f"| {sequence.get('month')} | {sequence.get('description')} | {sequence.get('record_count')} | "
            f"{fmt_money(sequence.get('total_amount'))} | {sequence.get('change_plan_id') or ''} | {sequence.get('source_file') or ''} |"
        )
    lines += ["", "## Nicht übertragene Vorgänge", ""]
    decided = {str(item.get("transaction_id")): item for item in transfer.get("not_transferred", [])}
    if not rejected and not decided:
        lines.append("Keine.")
    else:
        lines += ["| Vorgang | Grund | Betrag | Entscheidung | Entscheider | Datum |", "|---|---|---|---|---|---|"]
        seen = set()
        for item in rejected:
            tid = str(item.get("transaction_id"))
            seen.add(tid)
            entry = decided.get(tid, {})
            amount = fmt_money(item.get("amount")) if item.get("amount") is not None else "unsicher"
            lines.append(f"| {tid} | {item.get('reason')} | {amount} | {entry.get('decision') or 'offen'} | {entry.get('decided_by') or ''} | {entry.get('decided_at') or ''} |")
        for tid, entry in decided.items():
            if tid not in seen:
                lines.append(f"| {tid} | {entry.get('reason')} | | {entry.get('decision') or ''} | {entry.get('decided_by') or ''} | {entry.get('decided_at') or ''} |")
    lines += ["", "## Abweichungen vom EXTF-Paket", ""]
    deviations = []
    for sequence in records.get("records_rot", []) + records.get("records_gruen", []):
        if int(sequence.get("not_transferred_count", 0)):
            deviations.append(
                f"- {sequence['source_file']}: {sequence['extf_rows']} Zeilen / {fmt_money(sequence['extf_total'])} EUR im EXTF, "
                f"übertragen {sequence['record_count']} Zeilen / {fmt_money(sequence['total_amount'])} EUR."
            )
    if records.get("records_rot"):
        deviations.append(f"- Rote Zeilen stehen auf dem Klärungskonto {clearing.get('account')} statt auf leeren Feldern; Belegdatum gefüllt, KLÄR-Buchungstext.")
    for item in transfer.get("removed_documents", []):
        deviations.append(f"- Beleg {item.get('guid')} aus {item.get('file')} entfernt: {item.get('reason')}")
    lines += deviations or ["Keine."]
    lines += ["", "## Dateien, die nicht mehr importiert werden dürfen", ""]
    files = transfer.get("transferred_files") or records.get("transferred_files_expected") or []
    lines += [f"- {name} (verschoben nach {LOG_FOLDER}/{REPLACED_FOLDER}/)" if name in moved else f"- {name}" for name in files] or ["Keine."]
    lines += ["", "## Sichtbarkeit in DATEV", ""]
    visibility = transfer.get("visibility_check") or {}
    if visibility:
        lines.append(f"{visibility.get('tool')}: {visibility.get('result')} ({visibility.get('checked_at')}). Eine nicht nachgewiesene Verbuchung gilt nicht als bestätigt.")
    else:
        lines.append("Nicht geprüft. Die Verbuchung gilt erst nach Sichtung in DATEV als nachgewiesen.")
    lines += ["", "## Offene Schritte", ""]
    steps = list(transfer.get("open_steps") or [])
    if not steps:
        steps = [
            "Reguläre und Avis-Belegtransfer-ZIPs in DATEV Unternehmen online hochladen; erst danach lösen die Beleglinks auf.",
            "Klärungsposten vom Klärungskonto auf die Zielkonten umbuchen; Klärungskonto je Periode auf Saldo 0 bringen; erst danach festschreiben.",
        ]
    lines += [f"- {step}" for step in steps]
    lines.append("")
    return "\n".join(lines)


def main() -> int:
    args = parse_args()
    try:
        package = args.package
        records = json.loads(args.records.read_text(encoding="utf-8"))
        transfer = json.loads(args.transfer.read_text(encoding="utf-8"))
        rejected_path = args.records.with_name("nicht_uebertragbar.json")
        rejected = json.loads(rejected_path.read_text(encoding="utf-8")) if rejected_path.is_file() else []
        clearing = transfer.get("clearing_account") or records.get("clearing_account") or {}
        if not clearing.get("account"):
            raise ValueError("Klärungskonto fehlt in riecken_transfer.json und riecken_records.json.")
        if not transfer.get("sequences"):
            transfer["sequences"] = records.get("preview", [])
        replaced = package / LOG_FOLDER / REPLACED_FOLDER
        replaced.mkdir(parents=True, exist_ok=True)

        workbook_path = find_workbook(package, args.workbook)
        backup = replaced / f"{workbook_path.stem}_vor_Riecken.xlsx"
        if not backup.exists():
            shutil.copy2(workbook_path, backup)
        workbook = load_workbook(backup)
        by_row, by_transaction, red_by_transaction = record_index(records)
        decisions = {str(item.get("transaction_id")): item for item in transfer.get("not_transferred", [])}
        rejected_by_tid = {str(item.get("transaction_id")): item for item in rejected}
        update_bookings_sheet(workbook["Buchungszeilen"], by_row, by_transaction, red_by_transaction, decisions, rejected_by_tid, clearing)
        update_review_sheet(workbook["Belegprüfung"], by_transaction, red_by_transaction, decisions, rejected_by_tid, clearing)
        update_summary_sheet(workbook["Übersicht"], transfer, records, decisions)
        update_guide_sheet(workbook["Anleitung"], clearing)
        build_transfer_sheet(workbook, transfer, rejected)
        workbook.save(workbook_path)

        moved: list[str] = []
        files = transfer.get("transferred_files") or records.get("transferred_files_expected") or []
        if records.get("master_data_file") and records["master_data_file"] not in files and any(
            plan.get("type") == "business_partner" for plan in transfer.get("change_plans", [])
        ):
            files = list(files) + [records["master_data_file"]]
        if not args.keep_extf:
            for name in files:
                source = package / DATEV_FOLDER / name
                if source.is_file():
                    shutil.move(str(source), str(replaced / name))
                    moved.append(name)
        transfer["transferred_files"] = list(files)

        status_path = package / STATUS_FILE
        status_path.write_text(status_markdown(package, transfer, records, rejected, moved), encoding="utf-8")
        manifest_path = package / LOG_FOLDER / "Laufmanifest.json"
        if manifest_path.is_file():
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            manifest["riecken_transfer"] = transfer
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({
            "workbook": str(workbook_path), "backup": str(backup), "status": str(status_path),
            "moved_extf": moved, "transfer_sheet": TRANSFER_SHEET,
        }, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:  # pragma: no cover - CLI-Fehlerpfad
        print(f"FEHLER: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
