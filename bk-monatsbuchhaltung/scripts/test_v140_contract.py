from __future__ import annotations

import copy
import hashlib
import json
import subprocess
import sys
import tempfile
from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from openpyxl import load_workbook

import build_package
import datev_io
import validate_package
from sharepoint_target import build_targets
from test_datev_contract import booking_document, evidence


PERIOD = "2025-12"
PACKAGE_DIR = f"12861_{PERIOD}"
COST_CONFIG = {
    "kost_system": 1,
    "kost1_required": True,
    "kost2_required": False,
    "kost1_allowed": {"1000": "Praxis", "2000": "Labor", "9999": "Sammelkostenstelle/-träger"},
    "kost2_allowed": {},
    "rules_source": "Synthetisches Testprofil, Abschnitt Kostenstellenregeln",
}
BATCH_CONFIG = {
    "separate_batches": {
        "eigenbelege": {
            "label": "Eigenbelege Labor",
            "required_kost1": "2000",
            "required_contra_account": "8400",
        }
    }
}


def expect_error(callable_, fragment: str, label: str) -> None:
    try:
        callable_()
    except ValueError as exc:
        assert fragment in str(exc), f"{label}: unerwartete Meldung {exc}"
    else:
        raise AssertionError(f"{label}: Fehler wurde nicht erkannt")


def make_run(root: Path, *, cost: bool = False, pflicht: bool = False, batches: bool = False) -> dict:
    profile = root / "12861.md"
    if not profile.exists():
        profile.write_text("# Synthetisches Testprofil\n", encoding="utf-8")
    run = {
        "beraternummer": 29098, "mandantennummer": 12861,
        "buchungsmonat": PERIOD, "wirtschaftsjahr_beginn": "2025-01-01",
        "sachkontenlaenge": 4, "sachkontenrahmen": "03", "waehrung": "EUR",
        "accounting_method": "EÜR", "kostenstellenpflicht": pflicht,
        "datev_connection_verified": True, "mandantenprofil_verified": True,
        "vat_config": {"sales_treatment": "steuerpflichtig", "input_tax_deduction": "voll", "default_domestic_input_treatment": "volle_vorsteuer"},
        "account_config": {"private_expense": "4655", "gwg": "0480", "asset_accounts": ["0480", "0500"]},
        "person_account_ranges": {"debitor": {"start": "10000", "end": "69999"}, "kreditor": {"start": "70000", "end": "99999"}},
        "mandantenprofil_evidence": evidence(profile, str(build_targets("12861")["profile_url"])),
    }
    if cost:
        config = copy.deepcopy(COST_CONFIG)
        config["kost1_required"] = pflicht
        run["cost_center_config"] = config
    if batches:
        run["batch_config"] = copy.deepcopy(BATCH_CONFIG)
    run["datev_live_evidence"] = {
        **{key: run[key] for key in ("beraternummer", "mandantennummer", "wirtschaftsjahr_beginn", "sachkontenlaenge", "sachkontenrahmen")},
        "source": "DATEV live", "retrieved_at": "2026-10-07T10:00:00+02:00",
        "validated_accounts": ["4900", "4655", "8400", "70001", "10001"], "validated_bu_keys": ["401", "900"],
        "highest_creditor_account": 70001, "highest_debtor_account": 10001,
        "used_person_accounts": [
            {"account": "70001", "account_type": "kreditor", "name": "DATEV Test GmbH"},
            {"account": "10001", "account_type": "debitor", "name": "Labor Kunde"},
        ],
        "master_data_checked": True, "master_data_records_found": 2,
        "prior_bookings_checked": True, "prior_booking_records_found": 0,
    }
    if cost:
        run["datev_live_evidence"]["cost_system_active"] = True
        run["datev_live_evidence"]["validated_cost_centers"] = ["1000", "2000", "9999"]
    return run


class Scenario:
    """Collects synthetic documents and writes a complete run JSON."""

    def __init__(self, root: Path, run: dict) -> None:
        self.root = root
        self.run = run
        self.docs: list[dict] = []
        self.sources: list[dict] = []
        self.mappings: list[dict] = []
        self.cases: list[dict] = []
        self.releases: list[dict] = []

    def add(self, light: str = "Grün", **overrides) -> dict:
        number = len(self.docs) + 1
        source = self.root / f"beleg{number}.txt"
        source.write_text(f"Synthetischer Beleg {number} {light}", encoding="utf-8")
        doc = booking_document(source, light)
        doc["transaction_id"] = f"V{number:04d}"
        doc["invoice_number"] = f"RE-{number}"
        doc["bookings"][0]["document_field_1"] = f"RE-{number}"
        doc["entity_assessment"] = {"legal_entity": "Test", "addressee": "Test", "relevance": "in_scope"}
        doc["duplicate_checks"] = {key: {"checked": True, "result": "no_hit", "reference": ""} for key in ("file_hash_current_upload", "logical_document_current_upload", "datev_live")}
        bookings = overrides.pop("bookings", None)
        doc.update(overrides)
        if bookings:
            doc["bookings"] = bookings
        if light == "Rot":
            doc["requires_clarification"] = True
            doc.setdefault("reason", "Offene fachliche Entscheidung.")
            self.cases.append({
                "case_id": f"K{number}", "transaction_ids": [doc["transaction_id"]],
                "topic": "Klärung", "facts": "Sichere Angaben exportiert.",
                "booking_risk": "Ohne Klärung kann die Buchung fachlich falsch sein.",
                "provisional_treatment": "Offenes Feld leer exportiert.",
                "recommendation": "Originalbeleg prüfen.", "decision_needed": "Offenes Feld vervollständigen.",
                "traffic_light": "Rot", "target": "Prüfungsdatei", "proposed_change": "Feld klären", "employee_result": "",
            })
        self.docs.append(doc)
        source_id = f"S{number}"
        self.sources.append({"source_id": source_id, "source_path": str(source), "size_bytes": source.stat().st_size, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "readability": "readable"})
        self.mappings.append({"transaction_id": doc["transaction_id"], "source_id": source_id, "role": "primary_invoice"})
        return doc

    def data(self) -> dict:
        return {
            "run": self.run,
            "scope": {"target_periods": [PERIOD], "job_mode": "belegbuchhaltung", "include_prior_periods": False, "include_future_periods": False},
            "source_files": self.sources, "transactions": self.docs, "transaction_sources": self.mappings,
            "clarification_cases": self.cases, "master_records": [], "accrual_releases": self.releases,
            "activity_report": {"datev_import_status": "Importpaket erstellt – noch nicht in DATEV importiert", "sources_used": ["Synthetische Testbelege"], "named_entities": []},
        }

    def write(self, name: str = "lauf.json") -> Path:
        path = self.root / name
        path.write_text(json.dumps(self.data(), ensure_ascii=False), encoding="utf-8")
        return path

    def build(self, name: str = "package") -> tuple[Path, dict, dict]:
        input_path = self.write(f"{name}.json")
        result = subprocess.run(
            [sys.executable, str(Path(build_package.__file__)), "--input", str(input_path), "--output", str(self.root / name)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
        )
        assert result.returncode == 0, result.stderr + result.stdout
        package = self.root / name / PACKAGE_DIR
        manifest = json.loads((package / "03_Technische_Protokolle" / "Laufmanifest.json").read_text(encoding="utf-8"))
        report = json.loads((package / "03_Technische_Protokolle" / "Validierungsbericht.json").read_text(encoding="utf-8"))
        assert report["valid"], report["errors"]
        return package, manifest, report

    def load(self, name: str) -> dict:
        return build_package.load_input(self.write(f"{name}.json"))

    def document_errors(self, name: str) -> list[str]:
        data = self.load(name)
        return build_package.validate_documents(data)


def csv_rows(path: Path) -> list[list[str]]:
    return [validate_package.split_extf(line) for line in path.read_text(encoding="cp1252").splitlines()[2:]]


def csv_header(path: Path) -> list[str]:
    return validate_package.split_extf(path.read_text(encoding="cp1252").splitlines()[0])


def extf_names(package: Path) -> list[str]:
    return sorted(path.name for path in (package / "01_DATEV_Import").glob("EXTF_*.csv"))


def run_validator(package: Path) -> dict:
    subprocess.run(
        [sys.executable, str(Path(validate_package.__file__)), "--package", str(package)],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    return json.loads((package / "03_Technische_Protokolle" / "Validierungsbericht.json").read_text(encoding="utf-8"))


def red_booking(**fields) -> dict:
    line = {
        "amount": "119.00", "debit_credit": "S", "account": "4900", "account_name": "Betriebsbedarf",
        "contra_account": "70001", "contra_account_name": "DATEV Test GmbH", "bu_key": "401",
        "document_field_1": "RE-X", "booking_text": "Betriebsbedarf DATEV Test", "open_fields": {},
    }
    line.update(fields)
    return line


def test_clarification_batches(root: Path) -> None:
    """Spezifikation Klärungsstapel: getrennte Dateien, Pflichtleerung, Sortierregel, Teilung."""
    scenario = Scenario(root, make_run(root))
    scenario.add("Grün")
    scenario.add("Grün")
    scenario.add("Rot", reason="Kontierung offen.", bookings=[red_booking(account=None, open_fields={"account": "Kontierung aus Beleg nicht erkennbar."})])
    duplicate = scenario.add("Rot", reason="Mögliche Dublette.")
    duplicate["prior_booking_check"] = {"checked": True, "result": "moegliche_dublette", "references": ["BU 2025-11/17"]}
    duplicate["duplicate_checks"]["datev_live"] = {"checked": True, "result": "possible_duplicate", "reference": "BU 2025-11/17"}
    scenario.releases.append({**scenario.docs[0]["bookings"][0], "accrual_id": "A1", "period": PERIOD, "booking_date": "2025-12-31", "document_field_1": "ARAP-A1"})
    package, manifest, report = scenario.build("klaerung")
    names = extf_names(package)
    assert names == [f"EXTF_Buchungsstapel_{PERIOD}.csv", f"EXTF_Klaerungsposten_{PERIOD}.csv"], names
    booking = package / "01_DATEV_Import" / names[0]
    clarification = package / "01_DATEV_Import" / names[1]
    assert csv_header(clarification)[16] == "Klärungsposten"
    assert csv_header(booking)[16] == "Buchungsstapel"
    assert csv_header(clarification)[14:16] == csv_header(booking)[14:16]
    assert csv_header(clarification)[2] == "21" and csv_header(clarification)[4] == "13" and csv_header(clarification)[20] == "0"
    booking_rows = csv_rows(booking)
    clarification_rows = csv_rows(clarification)
    assert len(booking_rows) == 3 and len(clarification_rows) == 2
    assert sorted(row[9] for row in booking_rows) == ["1512", "1512", "3112"]
    assert all(row[9] == "" for row in clarification_rows), "Belegdatum muss im Klärungsstapel immer leer sein"
    assert any(row[10] == "ARAP-A1" for row in booking_rows), "Abgrenzungsauflösung gehört in den Buchungsstapel"
    # Sortierregel: leeres Konto (V0003) steht vor dem vollständigen Dublettenverdacht (V0004).
    assert clarification_rows[0][6] == "" and clarification_rows[1][6] == "4900"
    assert not datev_io.carry_order_violations(clarification_rows)
    trace = {item["transaction_id"]: item for item in manifest["booking_trace"]}
    assert trace["V0001"]["file"] == names[0] and trace["V0001"]["batch_kind"] == "buchung"
    assert trace["V0003"]["file"] == names[1] and trace["V0003"]["batch_kind"] == "klaerung"
    assert trace["V0003"]["carry_order_ok"] is True and trace["V0003"]["csv_row"] == 3
    assert trace["V0004"]["recognized_date"] == "2025-12-15"
    assert trace["ABGRENZUNG-A1-1"]["file"] == names[0]
    assert manifest["datev_import_order"] == datev_io.DATEV_IMPORT_ORDER
    assert [batch["file"] for batch in manifest["booking_batches"]] == names
    assert manifest["booking_batches"][1]["rows"] == 2 and manifest["batch_split_reasons"] == []
    assert report["booking_batches"][0]["rows"] == 3

    workbook = load_workbook(next((package / "02_Buchungspruefung").glob("*.xlsx")))
    review_rows = list(workbook["Belegprüfung"].rows)
    headers = [cell.value for cell in review_rows[0]]
    assert headers[3] == "Belegdatum laut Beleg" and "KOST1" not in headers
    by_id = {row[2].value: row for row in review_rows[1:]}
    assert by_id["V0001"][1].value == names[0] and by_id["V0003"][1].value == names[1]
    assert by_id["V0003"][3].value is not None, "sicher erkanntes Datum bleibt in der Prüfungsdatei"
    booking_sheet = list(workbook["Buchungszeilen"].rows)
    assert {row[1].value for row in booking_sheet[1:] if row[0].value == "Rot"} == {names[1]}
    overview = " ".join(str(cell.value) for row in workbook["Übersicht"] for cell in row)
    assert names[1] in overview and "DATEV-Stapel" in overview
    guide = " ".join(str(cell.value) for row in workbook["Anleitung"] for cell in row)
    assert "Klärungsstapel" in guide and "eigener Importvorgang" in guide
    workbook.close()
    protocol = (package / "03_Technische_Protokolle" / "Laufprotokoll.md").read_text(encoding="utf-8")
    activity = (package / "02_Buchungspruefung" / "Taetigkeitsnachweis.md").read_text(encoding="utf-8")
    assert names[1] in protocol and names[1] in activity

    # Testimportnachweis: alle erzeugten Dateien und carry_over_result sind Pflicht.
    files = [booking, clarification]
    tested = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in files}
    base = {"status": "confirmed", "tested_at": "2026-10-08", "datev_version": "DATEV Rechnungswesen 2026", "test_client": "Testbestand 99999", "evidence_reference": "Protokoll 1", "result_detail": "übernommen"}
    assert validate_package.validate_test_import({**base, "tested_files": {names[0]: tested[names[0]]}}, files)
    assert validate_package.validate_test_import({**base, "tested_files": tested}, files)
    carry = {field: "not_tested" for field in datev_io.CARRY_FIELDS}
    carry["account"] = "carried"
    assert validate_package.validate_test_import({**base, "tested_files": tested, "carry_over_result": carry}, files) == []
    assert validate_package.validate_test_import({**base, "tested_files": tested, "carry_over_result": {**carry, "bu_key": "maybe"}}, files)

    # Negativtests am erzeugten Paket.
    report = run_validator(package)
    assert report["valid"]
    tampered = copy.deepcopy(manifest)
    next(item for item in tampered["booking_trace"] if item["transaction_id"] == "V0001")["traffic_light"] = "Rot"
    errors, _ = validate_package.validate_csv(booking, tampered)
    assert any("rote Zeile im Buchungsstapel" in item for item in errors), errors
    tampered = copy.deepcopy(manifest)
    next(item for item in tampered["booking_trace"] if item["transaction_id"] == "V0003")["traffic_light"] = "Grün"
    errors, _ = validate_package.validate_csv(clarification, tampered)
    assert any("grüne Zeile im Klärungsstapel" in item for item in errors), errors
    tampered = copy.deepcopy(manifest)
    next(item for item in tampered["booking_trace"] if item["transaction_id"] == "V0003")["kind"] = "accrual"
    errors, _ = validate_package.validate_csv(clarification, tampered)
    assert any("Abgrenzungsauflösung gehört in den Buchungsstapel" in item for item in errors), errors
    tampered = copy.deepcopy(manifest)
    next(item for item in tampered["booking_trace"] if item["transaction_id"] == "ABGRENZUNG-A1-1")["file"] = f"EXTF_Buchungsstapel_{PERIOD}_02.csv"
    layout_errors = validate_package.validate_batch_files(package, tampered)
    assert any("Exportnachweis ohne EXTF-Datei" in item for item in layout_errors), layout_errors
    # Klärungsstapel mit gefülltem Datum und falscher Reihenfolge.
    lines = clarification.read_text(encoding="cp1252").splitlines()
    swapped = clarification.with_name("swapped.csv")
    swapped.write_text("\r\n".join([lines[0], lines[1], lines[3], lines[2]]) + "\r\n", encoding="cp1252")
    swapped_named = root / "swap" / f"EXTF_Klaerungsposten_{PERIOD}.csv"
    swapped_named.parent.mkdir(parents=True, exist_ok=True)
    swapped_named.write_bytes(swapped.read_bytes())
    errors, _ = validate_package.validate_csv(swapped_named, {"run_contract": manifest["run_contract"]})
    assert any("Schleppschutz verletzt" in item and "Konto" in item for item in errors), errors
    dated = root / "dated" / f"EXTF_Klaerungsposten_{PERIOD}.csv"
    dated.parent.mkdir(parents=True, exist_ok=True)
    dated.write_bytes(booking.read_bytes())
    errors, _ = validate_package.validate_csv(dated, {"run_contract": manifest["run_contract"]})
    assert any("Pflichtleerung" in item for item in errors), errors
    assert any("Stapelbezeichnung ist nicht Klärungsposten" in item for item in errors), errors
    # Ampelfarbe im Dateinamen, zweiter Buchungsstapel, leere Datei.
    for bad_name in (f"EXTF_Buchungsstapel_Rot_{PERIOD}.csv", f"EXTF_Buchungsstapel_{PERIOD}_02.csv", f"EXTF_Klaerungsposten_{PERIOD}_Unbekannt.csv"):
        bad_root = root / "bad" / bad_name.replace(".csv", "")
        (bad_root / "01_DATEV_Import").mkdir(parents=True)
        (bad_root / "01_DATEV_Import" / bad_name).write_bytes(booking.read_bytes())
        layout = validate_package.validate_datev_import_layout(bad_root, manifest)
        assert layout, bad_name
    empty_root = root / "empty"
    (empty_root / "01_DATEV_Import").mkdir(parents=True)
    empty_file = empty_root / "01_DATEV_Import" / f"EXTF_Klaerungsposten_{PERIOD}.csv"
    empty_file.write_text("\r\n".join(lines[:2]) + "\r\n", encoding="cp1252")
    assert any("leere EXTF-Datei" in item for item in validate_package.validate_batch_files(empty_root, None))
    gap_root = root / "gap"
    (gap_root / "01_DATEV_Import").mkdir(parents=True)
    (gap_root / "01_DATEV_Import" / f"EXTF_Klaerungsposten_{PERIOD}.csv").write_bytes(clarification.read_bytes())
    (gap_root / "01_DATEV_Import" / f"EXTF_Klaerungsposten_{PERIOD}_03.csv").write_bytes(clarification.read_bytes())
    assert any("lückenlos" in item for item in validate_package.validate_batch_files(gap_root, None))
    unneeded_root = root / "unneeded"
    (unneeded_root / "01_DATEV_Import").mkdir(parents=True)
    (unneeded_root / "01_DATEV_Import" / f"EXTF_Klaerungsposten_{PERIOD}.csv").write_bytes(clarification.read_bytes())
    (unneeded_root / "01_DATEV_Import" / f"EXTF_Klaerungsposten_{PERIOD}_02.csv").write_bytes(clarification.read_bytes())
    assert any("Teilung" in item and "unzulässig" in item for item in validate_package.validate_batch_files(unneeded_root, None))


def test_batch_presence(root: Path) -> None:
    """Monat ohne Rot: kein Klärungsstapel. Monat nur mit Rot: kein Buchungsstapel."""
    (root / "green").mkdir()
    green = Scenario(root / "green", make_run(root / "green"))
    green.add("Grün")
    package, manifest, _ = green.build("green")
    assert extf_names(package) == [f"EXTF_Buchungsstapel_{PERIOD}.csv"]
    # Regression ohne Kostenstellen: Exportwerte entsprechen dem v1.3-Zeilenvertrag, KOST-Felder leer.
    row = csv_rows(package / "01_DATEV_Import" / f"EXTF_Buchungsstapel_{PERIOD}.csv")[0]
    assert row[:11] == ["119,00", "S", "EUR", "", "", "", "4900", "70001", "0401", "1512", "RE-1"]
    assert row[36:39] == ["", "", ""]
    workbook = load_workbook(next((package / "02_Buchungspruefung").glob("*.xlsx")))
    assert "KOST1" not in [cell.value for cell in list(workbook["Belegprüfung"].rows)[0]]
    workbook.close()

    (root / "red").mkdir()
    red = Scenario(root / "red", make_run(root / "red"))
    red.add("Rot", reason="Betrag unklar.", bookings=[red_booking(amount=None, open_fields={"amount": "Betrag auf dem Beleg nicht lesbar."})], total_amount=None)
    package, manifest, _ = red.build("red")
    assert extf_names(package) == [f"EXTF_Klaerungsposten_{PERIOD}.csv"]
    assert manifest["booking_batches"][0]["batch_kind"] == "klaerung"


def test_split(root: Path) -> None:
    """Konfliktfall: Konto leer/Belegfeld 1 gefüllt gegen Konto gefüllt/Belegfeld 1 leer erzeugt _02."""
    # Hinweis: Das Belegdatum ist im Klärungsstapel immer leer (Pflichtleerung) und
    # kann deshalb keinen Konflikt mehr auslösen; der Konflikt wird über Belegfeld 1 erzeugt.
    scenario = Scenario(root, make_run(root))
    scenario.add("Rot", reason="Kontierung offen.", bookings=[red_booking(account=None, open_fields={"account": "Kontierung unklar."})])
    scenario.add("Rot", reason="Referenz offen.", invoice_number=None, bookings=[red_booking(document_field_1=None, open_fields={"document_field_1": "Keine Rechnungsnummer erkennbar."})])
    scenario.add("Rot", reason="Dublette.", bookings=[red_booking()])
    package, manifest, _ = scenario.build("split")
    names = extf_names(package)
    assert names == [f"EXTF_Klaerungsposten_{PERIOD}.csv", f"EXTF_Klaerungsposten_{PERIOD}_02.csv"], names
    for name in names:
        rows = csv_rows(package / "01_DATEV_Import" / name)
        assert not datev_io.carry_order_violations(rows), name
    assert len(manifest["batch_split_reasons"]) == 1 and "Teilung" in (package / "03_Technische_Protokolle" / "Laufprotokoll.md").read_text(encoding="utf-8")
    parts = {batch["file"]: batch["part"] for batch in manifest["booking_batches"]}
    assert parts == {names[0]: 1, names[1]: 2}
    assert run_validator(package)["valid"]


def test_cost_centers(root: Path) -> None:
    """Änderungsauftrag Kostenstellen Teil A."""
    # Pflicht ohne cost_center_config: Preflight-Abbruch.
    (root / "nocfg").mkdir()
    bare = Scenario(root / "nocfg", make_run(root / "nocfg", pflicht=True))
    bare.add("Grün")
    expect_error(lambda: bare.load("nocfg"), "unkonfigurierte Pflichtkostenstelle", "Pflicht ohne Konfiguration")
    # kostenstellenpflicht muss boolesch sein.
    (root / "nobool").mkdir()
    nobool = Scenario(root / "nobool", make_run(root / "nobool"))
    nobool.run["kostenstellenpflicht"] = "nein"
    nobool.add("Grün")
    expect_error(lambda: nobool.load("nobool"), "true oder false", "kostenstellenpflicht nicht boolesch")
    # Kostenstellen ohne Konfiguration sind unzulässig.
    (root / "nocost").mkdir()
    nocost = Scenario(root / "nocost", make_run(root / "nocost"))
    nocost.add("Grün", bookings=[red_booking(kost1="1000", open_fields={})])
    assert any("ohne cost_center_config" in item for item in nocost.document_errors("nocost"))

    # Kostenstellen eingerichtet, keine Pflicht: ableitbare KOST1 exportiert, sonst Grün mit leerem Feld 37.
    (root / "optional").mkdir()
    optional = Scenario(root / "optional", make_run(root / "optional", cost=True))
    optional.add("Grün", bookings=[red_booking(kost1="1000", open_fields={})])
    optional.add("Grün")
    package, manifest, _ = optional.build("optional")
    rows = csv_rows(package / "01_DATEV_Import" / f"EXTF_Buchungsstapel_{PERIOD}.csv")
    assert sorted(row[36] for row in rows) == ["", "1000"]
    workbook = load_workbook(next((package / "02_Buchungspruefung").glob("*.xlsx")))
    headers = [cell.value for cell in list(workbook["Belegprüfung"].rows)[0]]
    assert headers[headers.index("Kontierung") + 1] == "KOST1"
    kost_cells = {row[2].value: row[headers.index("KOST1")].value for row in list(workbook["Belegprüfung"].rows)[1:]}
    assert kost_cells == {"V0001": "1000", "V0002": "ohne Kostenstelle"}
    booking_headers = [cell.value for cell in list(workbook["Buchungszeilen"].rows)[0]]
    assert booking_headers[booking_headers.index("BU-Schlüssel") + 1] == "KOST1" and "KOST2" not in booking_headers
    assert "Kostenstellen" in " ".join(str(cell.value) for row in workbook["Anleitung"] for cell in row)
    workbook.close()
    # Validator: Pflicht ohne KOST1 im Buchungsstapel wird erkannt.
    strict = copy.deepcopy(manifest)
    strict["run_contract"]["kostenstellenpflicht"] = True
    errors, _ = validate_package.validate_csv(package / "01_DATEV_Import" / f"EXTF_Buchungsstapel_{PERIOD}.csv", strict)
    assert any("KOST1 fehlt bei Kostenstellenpflicht" in item for item in errors), errors

    # Pflicht: Grün mit gültiger KOST1, Rot mit offenem kost1, gemischte Rechnung.
    (root / "pflicht").mkdir()
    strict_run = make_run(root / "pflicht", cost=True, pflicht=True)
    pflicht = Scenario(root / "pflicht", strict_run)
    pflicht.add("Grün", bookings=[red_booking(kost1="1000", open_fields={})])
    pflicht.add("Rot", reason="Zuordnung Praxis/Labor unklar.", bookings=[red_booking(kost1=None, open_fields={"kost1": "Zuordnung Praxis/Labor aus Beleg nicht erkennbar."})])
    pflicht.add("Grün", total_amount="238.00", derivation="Gemischte Rechnung je Kostenstelle aufgeteilt.", bookings=[
        red_booking(kost1="1000", open_fields={}),
        red_booking(kost1="2000", open_fields={}),
    ])
    package, manifest, report = pflicht.build("pflicht")
    names = extf_names(package)
    assert names == [f"EXTF_Buchungsstapel_{PERIOD}.csv", f"EXTF_Klaerungsposten_{PERIOD}.csv"]
    booking_rows = csv_rows(package / "01_DATEV_Import" / names[0])
    assert sorted(row[36] for row in booking_rows) == ["1000", "1000", "2000"]
    clarification_rows = csv_rows(package / "01_DATEV_Import" / names[1])
    assert clarification_rows[0][36] == "" and report["valid"]
    assert report["intentional_open_fields"][0]["open_fields"] == {"kost1": "Zuordnung Praxis/Labor aus Beleg nicht erkennbar."}
    assert manifest["run_contract"]["cost_center_config"]["kost1_required"] is True
    assert manifest["preflight_evidence"]["datev"]["validated_cost_centers"] == ["1000", "2000", "9999"]
    workbook = load_workbook(next((package / "02_Buchungspruefung").glob("*.xlsx")))
    headers = [cell.value for cell in list(workbook["Belegprüfung"].rows)[0]]
    kost_cells = {row[2].value: row[headers.index("KOST1")].value for row in list(workbook["Belegprüfung"].rows)[1:]}
    assert kost_cells == {"V0001": "1000", "V0002": "offen", "V0003": "1000, 2000"}
    open_column = headers.index("Offener Punkt / nächster Schritt")
    assert "kost1" in str({row[2].value: row[open_column].value for row in list(workbook["Belegprüfung"].rows)[1:]}["V0002"])
    workbook.close()

    # Grün ohne KOST1 bei Pflicht: Generatorfehler.
    (root / "missing").mkdir()
    missing = Scenario(root / "missing", make_run(root / "missing", cost=True, pflicht=True))
    missing.add("Grün")
    assert any("kost1 fehlt ohne dokumentierte Unsicherheit" in item for item in missing.document_errors("missing"))
    # Unbekannte KOST1: Generatorfehler.
    (root / "unknown").mkdir()
    unknown = Scenario(root / "unknown", make_run(root / "unknown", cost=True, pflicht=True))
    unknown.add("Grün", bookings=[red_booking(kost1="3000", open_fields={})])
    unknown.run["datev_live_evidence"]["validated_cost_centers"].append("3000")
    assert any("nicht in kost1_allowed" in item for item in unknown.document_errors("unknown"))
    # KOST1 nicht live in DATEV: nur Rot mit offenem kost1 zulässig.
    (root / "live").mkdir()
    live = Scenario(root / "live", make_run(root / "live", cost=True, pflicht=True))
    live.add("Grün", bookings=[red_booking(kost1="9999", open_fields={})])
    live.run["datev_live_evidence"]["validated_cost_centers"] = ["1000", "2000"]
    expect_error(lambda: live.load("live"), "nicht live in DATEV vorhanden", "KOST1 ohne Livenachweis")
    live.docs[0]["traffic_light"] = "Rot"
    live.docs[0]["requires_clarification"] = True
    live.docs[0]["bookings"][0].update({"kost1": None, "open_fields": {"kost1": "Kostenstelle 9999 in DATEV anlegen."}})
    live.cases.append({"case_id": "K1", "transaction_ids": ["V0001"], "topic": "Kostenstelle in DATEV anlegen", "facts": "9999 fehlt live.", "booking_risk": "Buchung ohne Kostenstelle.", "provisional_treatment": "kost1 offen.", "recommendation": "Kostenstelle 9999 anlegen.", "decision_needed": "Anlegen.", "traffic_light": "Rot", "target": "DATEV", "proposed_change": "Kostenstelle anlegen", "employee_result": ""})
    assert live.document_errors("live-red") == []
    # Fehlender Livenachweis bei konfigurierten Kostenstellen.
    (root / "nolive").mkdir()
    nolive = Scenario(root / "nolive", make_run(root / "nolive", cost=True))
    nolive.run["datev_live_evidence"].pop("validated_cost_centers")
    nolive.add("Grün")
    expect_error(lambda: nolive.load("nolive"), "validated_cost_centers", "Livenachweis Kostenstellen")
    # Manifestprüfung des Validators.
    broken = copy.deepcopy(manifest)
    broken["run_contract"]["cost_center_config"] = None
    assert any("Unkonfigurierte Pflichtkostenstelle" in item for item in validate_package._validate_preflight_manifest(broken))


def test_separate_batches(root: Path) -> None:
    """Änderungsauftrag Kostenstellen Teil B: getrennter Vorlauf für Eigenbelege."""
    scenario = Scenario(root, make_run(root, cost=True, pflicht=True, batches=True))
    scenario.add("Grün", bookings=[red_booking(kost1="1000", open_fields={})])
    scenario.add("Grün", batch_type="eigenbelege", partner="Labor Kunde", document_type="Ausgangsrechnung", bookings=[
        red_booking(kost1="2000", account="10001", account_name="Labor Kunde", contra_account="8400", contra_account_name="Erlöse", bu_key="", open_fields={}),
    ], input_tax_treatment="keine_vorsteuer")
    package, manifest, report = scenario.build("eigen")
    names = extf_names(package)
    assert names == [f"EXTF_Buchungsstapel_{PERIOD}.csv", f"EXTF_Buchungsstapel_{PERIOD}_Eigenbelege.csv"], names
    separate = package / "01_DATEV_Import" / names[1]
    assert csv_header(separate)[16] == "Eigenbelege Labor"
    rows = csv_rows(separate)
    assert len(rows) == 1 and rows[0][36] == "2000" and rows[0][7] == "8400"
    trace = {item["transaction_id"]: item for item in manifest["booking_trace"]}
    assert trace["V0002"]["file"] == names[1] and trace["V0002"]["batch_type"] == "eigenbelege"
    assert manifest["run_contract"]["batch_config"] == BATCH_CONFIG
    workbook = load_workbook(next((package / "02_Buchungspruefung").glob("*.xlsx")))
    assert {row[2].value: row[1].value for row in list(workbook["Belegprüfung"].rows)[1:]} == {"V0001": names[0], "V0002": names[1]}
    workbook.close()
    assert run_validator(package)["valid"]

    # Falsche KOST1 oder falsches Gegenkonto im Eigenbeleg-Stapel: Generatorfehler.
    wrong = copy.deepcopy(scenario)
    wrong.docs[1]["bookings"][0]["kost1"] = "1000"
    assert any("verlangt KOST1 2000" in item for item in wrong.document_errors("wrong-kost"))
    wrong = copy.deepcopy(scenario)
    wrong.docs[1]["bookings"][0].update({"contra_account": "4900", "contra_account_name": "Betriebsbedarf"})
    assert any("verlangt Gegenkonto 8400" in item for item in wrong.document_errors("wrong-contra"))
    # Eigenbeleg-Vorgang ohne batch_config: Generatorfehler.
    without = copy.deepcopy(scenario)
    without.run.pop("batch_config")
    assert any("ohne entsprechende batch_config" in item for item in without.document_errors("no-batch-config"))
    # Validator: Suffix nur aus batch_config, Stapelbezeichnung aus dem Profil.
    stripped = copy.deepcopy(manifest)
    stripped["run_contract"]["batch_config"] = None
    assert any("stammt nicht aus batch_config" in item for item in validate_package.validate_datev_import_layout(package, stripped))
    errors, _ = validate_package.validate_csv(separate, stripped)
    assert any("nicht in batch_config konfiguriert" in item for item in errors), errors


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="bk_v140_") as name:
        root = Path(name)
        for case in (test_clarification_batches, test_batch_presence, test_split, test_cost_centers, test_separate_batches):
            case_root = root / case.__name__
            case_root.mkdir()
            case(case_root)
    print("v1.4 clarification batches, cost centers and separate batches: OK")


if __name__ == "__main__":
    main()
