from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REVIEW_SHEET = "Belegprüfung"
SUMMARY_SHEET = "Übersicht"
EDITABLE_HEADERS = {"Bearbeitungsstatus", "Mitarbeiter-Ergebnis"}
REQUIRED_HEADERS = [
    "Ampel-Einstufung",
    "Buchungsstapel",
    "Vorgangs-ID",
    "Belegdatum",
    "Geschäftspartner",
    "Belegfeld 1",
    "Betrag",
    "Währung",
    "Buchungsperiode",
    "Kontierung",
    "Ableitung",
    "Prüfergebnis / Ampelbegründung",
    "Offener Punkt / nächster Schritt",
    "Bearbeitungsstatus",
    "Mitarbeiter-Ergebnis",
]
CLOSING_STATUSES = {
    "unverändert übernommen": "unverändert übernommen",
    "geändert": "geändert",
    "nicht übernommen": "nicht übernommen",
}
OPEN_STATUS = "offen"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Ausgefülltes Prüfprotokoll mit der Ausgangsdatei vergleichen"
    )
    parser.add_argument("--original", required=True, type=Path)
    parser.add_argument("--returned", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--mandantenprofil", type=Path)
    parser.add_argument("--abgrenzungsregister", type=Path)
    return parser.parse_args()


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def canonical(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        return round(value, 10)
    return value


def normalize_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9äöüß]+", " ", clean(value).casefold()).strip()


def status_value(value: Any) -> str:
    return clean(value).casefold()


def find_header_row(ws, headers: list[str]) -> int:
    expected = set(headers)
    for row_number in range(1, min(ws.max_row, 25) + 1):
        values = {clean(cell.value) for cell in ws[row_number] if clean(cell.value)}
        if expected.issubset(values):
            return row_number
    raise ValueError(f"{ws.title}: Kopfzeile nicht gefunden")


def table_rows(ws, headers: list[str]) -> tuple[list[dict[str, Any]], int]:
    header_row = find_header_row(ws, headers)
    positions = {
        clean(cell.value): cell.column
        for cell in ws[header_row]
        if clean(cell.value) in headers
    }
    rows: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        row = {
            header: ws.cell(row_number, positions[header]).value
            for header in headers
        }
        if any(value not in (None, "") for value in row.values()):
            row["__row__"] = row_number
            rows.append(row)
    return rows, header_row


def sheet_matrix(ws) -> list[list[Any]]:
    result: list[list[Any]] = []
    for row in ws.iter_rows():
        values = [canonical(cell.value) for cell in row]
        while values and values[-1] is None:
            values.pop()
        result.append(values)
    while result and not result[-1]:
        result.pop()
    return result


def workbook_context(wb) -> tuple[str, str]:
    if SUMMARY_SHEET not in wb.sheetnames:
        raise ValueError(f"Blatt {SUMMARY_SHEET!r} fehlt")
    title = clean(wb[SUMMARY_SHEET]["A1"].value)
    match = re.fullmatch(r"Übersicht\s+(.+?)\s+[–-]\s+(\d{4}-\d{2})", title)
    if not match:
        raise ValueError(
            "Mandant und Periode konnten aus Übersicht!A1 nicht gelesen werden"
        )
    return match.group(1), match.group(2)


def read_optional(path: Path | None) -> str:
    if path is None:
        return ""
    return path.read_text(encoding="utf-8")


def reference_contains(reference: str, row: dict[str, Any]) -> bool:
    if not reference:
        return False
    normalized_reference = normalize_text(reference)
    transaction_id = normalize_text(row.get("Vorgangs-ID"))
    partner = normalize_text(row.get("Geschäftspartner"))
    if transaction_id and transaction_id in normalized_reference:
        return True
    partner_tokens = [token for token in partner.split() if len(token) >= 4]
    return bool(partner_tokens) and all(
        token in normalized_reference for token in partner_tokens
    )


def classify(row: dict[str, Any], status: str) -> str:
    if status not in CLOSING_STATUSES:
        return "weiterhin offener Klärungsfall"
    text = normalize_text(
        " ".join(
            clean(row.get(header))
            for header in (
                "Mitarbeiter-Ergebnis",
                "Offener Punkt / nächster Schritt",
                "Prüfergebnis / Ampelbegründung",
                "Ableitung",
                "Kontierung",
            )
        )
    )
    if any(word in text for word in ("arap", "prap", "abgrenz", "auflösung")):
        return "Änderung oder Ergänzung einer Rechnungsabgrenzung"
    if any(
        word in text
        for word in ("personenkonto", "debitor", "kreditor", "lieferantenkonto")
    ):
        return "Bestätigung oder Änderung von Personenkonten"
    if any(
        word in text
        for word in ("dauerhaft", "künftig", "zukünftig", "immer", "mandantenregel")
    ):
        return "dauerhaft wiederverwendbare Mandantenbesonderheit"
    return "einmalige Buchungskorrektur"


def output_name(mandant: str, period: str) -> str:
    safe_mandant = re.sub(r'[<>:"/\\|?*]+', "_", mandant).strip(" .")
    safe_period = re.sub(r'[<>:"/\\|?*]+', "_", period).strip(" .")
    return f"Rücklaufauswertung {safe_mandant} {safe_period}.md"


def bullet(items: list[str], empty: str = "Keine.") -> list[str]:
    return [f"- {item}" for item in items] if items else [empty]


def evaluate(
    original_path: Path,
    returned_path: Path,
    output_dir: Path,
    mandantenprofil_path: Path | None = None,
    abgrenzungsregister_path: Path | None = None,
) -> dict[str, Any]:
    original = load_workbook(original_path, data_only=True)
    returned = load_workbook(returned_path, data_only=True)
    integrity: list[str] = []
    validation: list[str] = []

    original_context = workbook_context(original)
    returned_context = workbook_context(returned)
    mandant, period = original_context
    if returned_context != original_context:
        integrity.append(
            "Mandant oder Buchungsperiode weicht von der Ausgangsdatei ab "
            f"({returned_context[0]}/{returned_context[1]} statt {mandant}/{period})."
        )

    if original.sheetnames != returned.sheetnames:
        integrity.append("Blattstruktur wurde gegenüber der Ausgangsdatei verändert.")
    for sheet_name in original.sheetnames:
        if sheet_name == REVIEW_SHEET or sheet_name not in returned.sheetnames:
            continue
        if sheet_matrix(original[sheet_name]) != sheet_matrix(returned[sheet_name]):
            integrity.append(
                f"Blatt {sheet_name!r} wurde außerhalb der Rücklauffelder verändert."
            )

    if REVIEW_SHEET not in original.sheetnames or REVIEW_SHEET not in returned.sheetnames:
        raise ValueError(f"Blatt {REVIEW_SHEET!r} fehlt")
    source_rows, _ = table_rows(original[REVIEW_SHEET], REQUIRED_HEADERS)
    return_rows, _ = table_rows(returned[REVIEW_SHEET], REQUIRED_HEADERS)
    for source_row in source_rows:
        if clean(source_row.get("Ampel-Einstufung")) not in {"Rot", "Grün", ""}:
            validation.append(f"{source_row.get('Vorgangs-ID')}: Ampelstatus ist in Version 1.3 unzulässig.")

    def index_rows(rows: list[dict[str, Any]], label: str) -> dict[str, dict[str, Any]]:
        index: dict[str, dict[str, Any]] = {}
        for row in rows:
            transaction_id = clean(row.get("Vorgangs-ID"))
            if not transaction_id:
                integrity.append(f"{label}: leere Vorgangs-ID in Zeile {row['__row__']}.")
            elif transaction_id in index:
                integrity.append(f"{label}: doppelte Vorgangs-ID {transaction_id}.")
            else:
                index[transaction_id] = row
        return index

    source_by_id = index_rows(source_rows, "Ausgangsdatei")
    return_by_id = index_rows(return_rows, "Rücklaufdatei")
    missing_ids = sorted(set(source_by_id) - set(return_by_id))
    extra_ids = sorted(set(return_by_id) - set(source_by_id))
    if missing_ids:
        integrity.append("Gelöschte Vorgangs-IDs: " + ", ".join(missing_ids) + ".")
    if extra_ids:
        integrity.append("Hinzugefügte Vorgangs-IDs: " + ", ".join(extra_ids) + ".")

    immutable_headers = [
        header for header in REQUIRED_HEADERS if header not in EDITABLE_HEADERS
    ]
    for transaction_id in sorted(set(source_by_id) & set(return_by_id)):
        source_row = source_by_id[transaction_id]
        return_row = return_by_id[transaction_id]
        changed = [
            header
            for header in immutable_headers
            if canonical(source_row.get(header)) != canonical(return_row.get(header))
        ]
        if changed:
            integrity.append(
                f"{transaction_id}: geschützte Felder verändert: {', '.join(changed)}."
            )

    required_rows = [
        return_by_id[transaction_id]
        for transaction_id, source_row in source_by_id.items()
        if clean(source_row.get("Ampel-Einstufung")) == "Rot"
        and transaction_id in return_by_id
    ]
    results: list[dict[str, Any]] = []
    open_ids: list[str] = []
    counts = {
        "Rot": {"gesamt": 0, "abgeschlossen": 0},
    }
    for source_row in source_rows:
        light = clean(source_row.get("Ampel-Einstufung"))
        if light in counts:
            counts[light]["gesamt"] += 1

    for row in required_rows:
        transaction_id = clean(row.get("Vorgangs-ID"))
        light = clean(row.get("Ampel-Einstufung"))
        raw_status = status_value(row.get("Bearbeitungsstatus"))
        employee_result = clean(row.get("Mitarbeiter-Ergebnis"))
        if raw_status in CLOSING_STATUSES:
            status = CLOSING_STATUSES[raw_status]
            counts[light]["abgeschlossen"] += 1
            if status in {"geändert", "nicht übernommen"} and not employee_result:
                validation.append(
                    f"{transaction_id}: Mitarbeiter-Ergebnis ist bei Status {status!r} Pflicht."
                )
        elif raw_status in {"", OPEN_STATUS}:
            status = OPEN_STATUS
            open_ids.append(transaction_id)
        else:
            status = raw_status
            validation.append(
                f"{transaction_id}: unzulässiger Bearbeitungsstatus {clean(row.get('Bearbeitungsstatus'))!r}."
            )
        results.append(
            {
                "transaction_id": transaction_id,
                "light": light,
                "partner": clean(row.get("Geschäftspartner")),
                "status": status,
                "employee_result": employee_result,
                "category": classify(row, status),
                "row": row,
            }
        )

    profile = read_optional(mandantenprofil_path)
    register = read_optional(abgrenzungsregister_path)
    profile_suggestions: list[str] = []
    register_suggestions: list[str] = []
    person_account_impacts: list[str] = []
    changed_items: list[str] = []
    rejected_items: list[str] = []
    for result in results:
        label = f"{result['transaction_id']} ({result['partner'] or 'ohne Partner'})"
        details = result["employee_result"] or "kein zusätzlicher Ergebnistext"
        if result["status"] == "geändert":
            changed_items.append(f"{label}: {details}")
        elif result["status"] == "nicht übernommen":
            rejected_items.append(f"{label}: {details}")
        if result["category"] == "dauerhaft wiederverwendbare Mandantenbesonderheit":
            if not reference_contains(profile, result["row"]):
                profile_suggestions.append(
                    f"{label}: {details} (nur nach ausdrücklicher Freigabe übernehmen)"
                )
        elif result["category"] == "Änderung oder Ergänzung einer Rechnungsabgrenzung":
            if not reference_contains(register, result["row"]):
                register_suggestions.append(
                    f"{label}: {details} (nur nach ausdrücklicher Freigabe übernehmen)"
                )
        elif result["category"] == "Bestätigung oder Änderung von Personenkonten":
            person_account_impacts.append(f"{label}: {details}")

    remaining = [*integrity, *validation]
    remaining.extend(f"{transaction_id}: Bearbeitungsstatus offen." for transaction_id in open_ids)
    complete = not remaining and all(
        values["gesamt"] == values["abgeschlossen"] for values in counts.values()
    )
    overall = "vollständig" if complete else "unvollständig"

    categories: dict[str, list[str]] = {}
    for result in results:
        categories.setdefault(result["category"], []).append(
            f"{result['transaction_id']} ({result['status']})"
        )

    lines = [
        f"# Rücklaufauswertung {mandant} {period}",
        "",
        f"- Mandant: {mandant}",
        f"- Periode: {period}",
        f"- Geprüfte Quelldatei: {original_path.name}",
        f"- Rücklaufdatei: {returned_path.name}",
        f"- Integrität: {'bestätigt' if not integrity else 'abweichend'}",
        f"- Gesamtstatus: {overall}",
        "",
        "## Abschlussstatus Rot",
        "",
        f"- Rot: {counts['Rot']['abgeschlossen']} von {counts['Rot']['gesamt']} abgeschlossen",
        "",
        "## Geänderte Vorgänge",
        "",
        *bullet(changed_items),
        "",
        "## Nicht übernommene Vorgänge",
        "",
        *bullet(rejected_items),
        "",
        "## Fachliche Kategorien",
        "",
    ]
    if categories:
        for category, items in sorted(categories.items()):
            lines.append(f"- {category}: {', '.join(items)}")
    else:
        lines.append("Keine.")
    lines.extend(
        [
            "",
            "## Auswirkungen auf Mandantenprofil",
            "",
            *bullet(profile_suggestions, "Keine Änderung vorgeschlagen."),
            "",
            "## Auswirkungen auf Abgrenzungsregister",
            "",
            *bullet(register_suggestions, "Keine Änderung vorgeschlagen."),
            "",
            "## Auswirkungen auf Personenkonten",
            "",
            *bullet(person_account_impacts, "Keine Änderung festgestellt."),
            "",
            "## Verbleibende offene Punkte",
            "",
            *bullet(remaining),
            "",
            "Mandantenprofil und Abgrenzungsregister wurden nicht verändert. Vorschläge dürfen erst nach ausdrücklicher Freigabe übernommen werden.",
            "Ein DATEV-Buchungsstapel wird durch diese Auswertung nicht automatisch neu erzeugt.",
            "",
        ]
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / output_name(mandant, period)
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return {
        "mandant": mandant,
        "period": period,
        "status": overall,
        "integrity": not integrity,
        "counts": counts,
        "errors": remaining,
        "profile_suggestions": profile_suggestions,
        "register_suggestions": register_suggestions,
        "output": str(output_path),
    }


def main() -> int:
    args = parse_args()
    try:
        result = evaluate(
            args.original,
            args.returned,
            args.output_dir,
            args.mandantenprofil,
            args.abgrenzungsregister,
        )
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({"status": "fehler", "error": str(exc)}, ensure_ascii=False))
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "vollständig" else 2


if __name__ == "__main__":
    raise SystemExit(main())
