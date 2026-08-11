#!/usr/bin/env python3
"""Create a German shareholder resolution from validated JSON input."""

from __future__ import annotations

import argparse
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def decimal_value(value: Any, field: str) -> Decimal:
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    if not isinstance(value, str):
        raise ValueError(f"{field} must be a number or string")
    cleaned = value.strip().replace("EUR", "").replace("€", "").replace(" ", "")
    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount for {field}: {value}") from exc


def euro(value: Decimal) -> str:
    quantized = value.copy_abs().quantize(Decimal("0.01"))
    integer, fraction = f"{quantized:.2f}".split(".")
    groups = []
    while integer:
        groups.insert(0, integer[-3:])
        integer = integer[:-3]
    return f"{'.'.join(groups)},{fraction}"


def require_text(data: dict[str, Any], key: str) -> str:
    value = data.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Missing required text field: {key}")
    return value.strip()


def distribution_sentence(appropriation: dict[str, Any]) -> str:
    settlement = appropriation.get("settlement", "payout")
    if settlement == "none":
        return ""
    if settlement == "payout":
        return " Der Ausschüttungsbetrag wird nach Abzug der einzubehaltenden Steuern an die Gesellschafter ausgezahlt."
    if settlement == "shareholder-account":
        return " Der Ausschüttungsbetrag wird nach Abzug der einzubehaltenden Steuern mit dem Gesellschafterkonto verrechnet."
    if settlement == "custom":
        text = appropriation.get("custom_settlement_text")
        if not isinstance(text, str) or not text.strip():
            raise ValueError("custom_settlement_text is required for settlement=custom")
        return " " + text.strip()
    raise ValueError(f"Unsupported settlement: {settlement}")


def appropriation_text(result: Decimal, appropriation: dict[str, Any]) -> str:
    mode = appropriation.get("mode")
    if mode == "carryforward":
        if result > 0:
            return f"Der Jahresüberschuss in Höhe von EUR {euro(result)} wird auf neue Rechnung vorgetragen."
        if result < 0:
            return f"Der Jahresfehlbetrag in Höhe von EUR {euro(result)} wird auf neue Rechnung vorgetragen."
        return "Eine Ergebnisverwendung ist aufgrund des ausgeglichenen Jahresergebnisses nicht erforderlich."

    if mode == "custom":
        text = appropriation.get("custom_text")
        if not isinstance(text, str) or not text.strip():
            raise ValueError("custom_text is required for mode=custom")
        return text.strip()

    amount = decimal_value(appropriation.get("distribution_amount"), "distribution_amount")
    due_date = appropriation.get("due_date")
    if amount <= 0:
        raise ValueError("distribution_amount must be positive")
    if not isinstance(due_date, str) or not due_date.strip():
        raise ValueError("due_date is required for distributions")
    due_date = due_date.strip()
    settlement = distribution_sentence(appropriation)

    if mode == "full-distribution":
        if result <= 0:
            raise ValueError("A full distribution from current earnings requires a positive annual result")
        if amount != result:
            raise ValueError("For full-distribution, distribution_amount must equal annual_result")
        return (
            f"Der Jahresüberschuss in Höhe von EUR {euro(result)} wird vollständig ausgeschüttet. "
            f"Der Ausschüttungsbetrag ist am {due_date} fällig.{settlement}"
        )

    if mode == "partial-distribution":
        if result <= 0:
            raise ValueError("A partial distribution from current earnings requires a positive annual result")
        if amount >= result:
            raise ValueError("For partial-distribution, distribution_amount must be less than annual_result")
        rest = result - amount
        return (
            f"Vom Jahresüberschuss in Höhe von EUR {euro(result)} werden EUR {euro(amount)} ausgeschüttet. "
            f"Der verbleibende Betrag von EUR {euro(rest)} wird auf neue Rechnung vorgetragen. "
            f"Der Ausschüttungsbetrag ist am {due_date} fällig.{settlement}"
        )

    if mode == "retained-earnings-distribution":
        if result > 0:
            current = f"Der Jahresüberschuss in Höhe von EUR {euro(result)} wird auf neue Rechnung vorgetragen."
        elif result < 0:
            current = f"Der Jahresfehlbetrag in Höhe von EUR {euro(result)} wird auf neue Rechnung vorgetragen."
        else:
            current = "Eine Ergebnisverwendung ist aufgrund des ausgeglichenen Jahresergebnisses nicht erforderlich."
        return (
            f"{current} Zusätzlich werden aus dem Gewinnvortrag EUR {euro(amount)} ausgeschüttet. "
            f"Der Ausschüttungsbetrag ist am {due_date} fällig.{settlement}"
        )

    raise ValueError(f"Unsupported appropriation mode: {mode}")


def set_font(run, size: float = 11, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Seite ")
    set_font(run, size=9, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_numbered_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_font(run)


def build_document(data: dict[str, Any], output: Path) -> None:
    company = require_text(data, "company_name")
    office = require_text(data, "registered_office")
    year = int(data.get("fiscal_year"))
    balance = decimal_value(data.get("balance_total"), "balance_total")
    result = decimal_value(data.get("annual_result"), "annual_result")
    if balance < 0:
        raise ValueError("balance_total must not be negative")
    appropriation = data.get("appropriation")
    if not isinstance(appropriation, dict):
        raise ValueError("appropriation must be an object")

    if result > 0:
        result_phrase = f"einem Jahresüberschuss von EUR {euro(result)}"
    elif result < 0:
        result_phrase = f"einem Jahresfehlbetrag von EUR {euro(result)}"
    else:
        result_phrase = "einem ausgeglichenen Jahresergebnis von EUR 0,00"

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.right_margin = Cm(2.5)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    list_style = doc.styles["List Number"]
    list_style.font.name = "Arial"
    list_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    list_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    list_style.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run(company)
    set_font(header_run, size=9, color="666666")

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(24)
    title_run = title.add_run(f"Gesellschafterversammlung der\n{company}, {office}")
    set_font(title_run, size=16, bold=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(16)
    intro.paragraph_format.line_spacing = 1.15
    intro_run = intro.add_run(
        f"Unter Verzicht auf alle Formen und Fristen halten die Gesellschafter der {company}, {office}, "
        "eine Gesellschafterversammlung ab und beschließen einstimmig, was folgt:"
    )
    set_font(intro_run)

    add_numbered_paragraph(
        doc,
        f"Der Jahresabschluss zum 31.12.{year} mit einer Bilanzsumme von EUR {euro(balance)} "
        f"und {result_phrase} wird festgestellt.",
    )

    if data.get("discharge_management", True):
        add_numbered_paragraph(doc, "Der Geschäftsführung wird für das abgelaufene Geschäftsjahr Entlastung erteilt.")

    add_numbered_paragraph(doc, appropriation_text(result, appropriation))

    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(10)
    closing.paragraph_format.space_after = Pt(28)
    closing_run = closing.add_run(
        "Weitere Beschlüsse werden nicht gefasst. Damit ist die Gesellschafterversammlung beendet."
    )
    set_font(closing_run)

    resolution_date = data.get("resolution_date")
    date_text = str(resolution_date).strip() if resolution_date else "____________________"
    date_para = doc.add_paragraph()
    date_para.paragraph_format.space_after = Pt(38)
    date_run = date_para.add_run(f"{office}, den {date_text}")
    set_font(date_run)

    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(0)
    line_run = line.add_run("________________________________________")
    set_font(line_run)

    signature_label = data.get("signature_label", "Unterschrift(en) der Gesellschafter")
    label = doc.add_paragraph()
    label.paragraph_format.space_after = Pt(0)
    label_run = label.add_run(str(signature_label))
    set_font(label_run)

    doc.core_properties.title = f"Gesellschafterbeschluss {company} {year}"
    doc.core_properties.subject = f"Feststellung des Jahresabschlusses zum 31. Dezember {year}"
    doc.core_properties.keywords = "Gesellschafterbeschluss, Jahresabschluss, Ergebnisverwendung"

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Path to validated JSON input")
    parser.add_argument("--output", required=True, type=Path, help="Path to the output DOCX")
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    build_document(data, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
