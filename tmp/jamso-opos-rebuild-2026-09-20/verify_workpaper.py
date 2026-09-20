from pathlib import Path
from decimal import Decimal
from collections import defaultdict
import json
import openpyxl

root=Path(__file__).parent
file=Path('outputs/01a0bea8-f4c0-74a2-b3e0-eafdb8a43b7c/jamso-opos-neuberechnung-2026-09-20.xlsx')
data=json.loads((root/'workpaper_data.json').read_text(encoding='utf8'))
wf=openpyxl.load_workbook(file,data_only=False)
wv=openpyxl.load_workbook(file,data_only=True)
def d(v): return Decimal(str(v or 0))
def cents(v): return d(v).quantize(Decimal('0.01'))
assert wf.sheetnames==['Arbeitspapier','Kontenabgleich','Nachtrag','Aufteilung 2024-25','Mandantenzeilen','DATEV 31.12.24','Zahlungstracker','Stornoliste']
assert len(data['datev_open_rows'])==789
assert len(data['later_rows'])==24
assert len(data['split_rows'])==26
assert len(data['tracker_rows'])==93
assert len(data['storno_rows'])==14
raw=wv['Mandantenzeilen']
for i,x in enumerate(data['source_rows'],5):
    assert cents(raw[f'L{i}'].value)==cents(d(x['debit'])-d(x['paid'])),('source now',i)
    expected=(d(x['debit'])-(d(x['paid']) if x['payment_year'] is not None and x['payment_year']<=2024 else 0)) if x['section']=='Hauptliste' else Decimal(0)
    assert cents(raw[f'M{i}'].value)==cents(expected),('source cutoff',i)
    assert wf['Mandantenzeilen'][f'L{i}'].data_type=='f'
datev=wv['DATEV 31.12.24']
for i,x in enumerate(data['datev_open_rows'],5):
    assert cents(datev[f'I{i}'].value)==cents(d(x['debit'])-d(x['credit'])),('datev row',i)
    assert wf['DATEV 31.12.24'][f'I{i}'].data_type=='f'
compare=wv['Kontenabgleich']
for i,x in enumerate(data['accounts'],5):
    a=x['account']
    source=[z for z in data['source_rows'] if z['section']=='Hauptliste' and z['account']==a]
    dv=[z for z in data['datev_open_rows'] if z['account']==a]
    expected_now=sum((d(z['debit'])-d(z['paid']) for z in source),Decimal(0))
    expected_2024=sum((d(z['debit'])-(d(z['paid']) if z['payment_year'] is not None and z['payment_year']<=2024 else 0) for z in source),Decimal(0))
    expected_datev=sum((d(z['debit'])-d(z['credit']) for z in dv),Decimal(0))
    assert cents(compare[f'C{i}'].value)==cents(expected_now),('account current',a)
    assert cents(compare[f'D{i}'].value)==cents(expected_2024),('account cutoff',a)
    assert cents(compare[f'E{i}'].value)==cents(expected_datev),('account datev',a)
    assert cents(compare[f'F{i}'].value)==cents(expected_2024-expected_datev),('account diff',a)
summary=wv['Arbeitspapier']
expected={8:'87114.51',9:'56387.07',10:'30727.44',11:'102675.71',12:'67466.20',13:'249441.37',14:'-1357.00'}
for r,v in expected.items(): assert cents(summary[f'B{r}'].value)==Decimal(v),('headline',r)
assert summary['B15'].value==69
assert wf['Mandantenzeilen']['J6'].number_format=='dd.mm.yyyy'
assert wf['Nachtrag']['E5'].number_format=='dd.mm.yyyy'
assert wf['DATEV 31.12.24']['C5'].number_format=='dd.mm.yyyy'
for i,x in enumerate(data['split_rows'],5):
    assert cents(wv['Aufteilung 2024-25'][f'H{i}'].value)==cents(d(x['part_2024'])+d(x['part_2025']))
    if x['tracker_invoice'] is not None:
        assert cents(wv['Aufteilung 2024-25'][f'K{i}'].value)==cents(d(x['part_2024'])+d(x['part_2025'])-d(x['tracker_invoice']))
print({'status':'PASS','source_rows_checked':len(data['source_rows']),'datev_rows_checked':len(data['datev_open_rows']),'account_rows_checked':len(data['accounts']),'split_rows_checked':len(data['split_rows']),'summary':expected})
