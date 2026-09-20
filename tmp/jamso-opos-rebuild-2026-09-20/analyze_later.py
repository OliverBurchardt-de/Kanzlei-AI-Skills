from pathlib import Path
from collections import defaultdict
from decimal import Decimal
import json
import re
import openpyxl

p=Path(__file__).parent
s=openpyxl.load_workbook(p/'debit.xlsx',data_only=False).active
q=json.loads((p/'datev_2024.json').read_text(encoding='utf8'))
norm=lambda x: str(x).replace('Debi','').replace(' ','')
for r in range(201,225):
    ref=s.cell(r,3).value
    name=s.cell(r,2).value
    pay=s.cell(r,5).value
    if not ref and not name: continue
    accts=defaultdict(lambda: Decimal(0))
    all_accts=set()
    for x in q:
        a=int(x['account_number'])//10000
        hit= ref and norm(ref) in norm(x.get('document_field1',''))+' '+norm(x.get('posting_description',''))
        if hit:
            all_accts.add(a)
            if not x['is_cleared']:
                accts[a]+=Decimal(str(x.get('amount_debit') or 0))-Decimal(str(x.get('amount_credit') or 0))
    print(r,repr(ref),repr(name),pay,'accts',sorted(all_accts),'open sums',dict(accts))
