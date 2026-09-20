from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl


ROOT = Path(__file__).parent
source = openpyxl.load_workbook(ROOT / 'debit.xlsx', data_only=False).active
datev = json.loads((ROOT / 'datev_2024.json').read_text(encoding='utf-8'))
split_sheet = openpyxl.load_workbook(ROOT / 'split.xlsx', data_only=False)['2024_2025']
tracker_file = openpyxl.load_workbook(ROOT / 'payments.xlsx', data_only=True)


def amount(value):
    return float(Decimal(str(value))) if isinstance(value, (int, float)) else 0.0


source_rows = []
source_accounts = {}
active_account = None
active_name = None
for r in list(range(2, 200)) + list(range(201, 225)):
    raw_account = source.cell(r, 1).value
    raw_name = source.cell(r, 2).value
    reference = source.cell(r, 3).value
    debit = source.cell(r, 4).value
    paid = source.cell(r, 5).value
    raw_date = source.cell(r, 6).value
    if not any(v is not None for v in (raw_account, raw_name, reference, debit, paid, raw_date)):
        continue
    section = 'Hauptliste' if r < 200 else 'Nachtrag'
    if section == 'Hauptliste':
        if isinstance(raw_account, int):
            active_account = raw_account
            active_name = str(raw_name or '')
            source_accounts[active_account] = active_name
        account = active_account
        name = str(raw_name or active_name or '')
    else:
        account = raw_account if isinstance(raw_account, int) else None
        name = str(raw_name or '')
    note = ''
    if isinstance(raw_date, datetime):
        parsed_date = raw_date.strftime('%Y-%m-%d')
    elif r == 50 and raw_date == '11.04.024':
        parsed_date = '2024-04-11'
        note = 'Datumsfehler in Quelle; DATEV zeigt Zahlung am 11.04.2024.'
    else:
        parsed_date = None
        if raw_date is not None:
            note = 'Datum nicht maschinenlesbar; nicht stillschweigend geändert.'
    if r == 77:
        note = 'Quelle nennt 05.02.2026, DATEV 05.02.2024; Jahr offen klären.'
    if r == 82:
        note = 'Quelle nennt 19.02.2023, DATEV 19.02.2024; Jahr offen klären.'
    if r == 170:
        note = 'Quelle: „neu gebucht in 2025“; DATEV zeigt 16.09.2024 offen.'
    if r == 224:
        note = 'Rechnung 2024-279 auch in Zeile 223, dort anderer Name.'
    source_rows.append({
        'source_row': r,
        'section': section,
        'raw_account': raw_account if isinstance(raw_account, int) else None,
        'account': account,
        'name': name,
        'reference': str(reference) if reference is not None else '',
        'debit': amount(debit),
        'paid': amount(paid),
        'raw_date': raw_date.strftime('%Y-%m-%d') if isinstance(raw_date, datetime) else str(raw_date or ''),
        'parsed_date': parsed_date,
        'payment_year': int(parsed_date[:4]) if parsed_date else None,
        'note': note,
    })

open_rows = []
datev_accounts = defaultdict(list)
for idx, obj in enumerate(datev, 1):
    if obj.get('is_cleared'):
        continue
    account = int(obj['account_number']) // 10000
    datev_accounts[account].append(obj)
    open_rows.append({
        'source_index': idx,
        'account': account,
        'date': obj['date'][:10],
        'document': str(obj.get('document_field1') or ''),
        'open_item': str(obj.get('open_item_number') or ''),
        'description': str(obj.get('posting_description') or ''),
        'debit': amount(obj.get('amount_debit')),
        'credit': amount(obj.get('amount_credit')),
        'id': str(obj.get('id') or ''),
    })

def clean_ref(ref):
    return re.sub(r'[^0-9a-z]', '', str(ref).lower()).replace('debi', '')

account_refs = defaultdict(set)
for obj in datev:
    account = int(obj['account_number']) // 10000
    for text in (obj.get('document_field1'), obj.get('posting_description')):
        text = str(text or '')
        for match in re.finditer(r'2024[- ]?\d{2,3}[a-z]?', text, re.I):
            account_refs[clean_ref(match.group())].add(account)

later_rows = []
last_ref = ''
last_name = ''
for item in source_rows:
    if item['section'] != 'Nachtrag':
        continue
    if item['reference']:
        last_ref = item['reference']
    if item['name']:
        last_name = item['name']
    ref = item['reference'] or last_ref
    name = item['name'] or last_name
    matches = sorted(account_refs.get(clean_ref(ref), [])) if ref else []
    if ref == '2024-279':
        match_status = 'Abweichender Zahlername: Reyes/Wothe'
        assigned = None
    elif len(matches) == 1:
        match_status = 'Rechnungsnr. eindeutig; Zahlung ungeprüft'
        assigned = matches[0]
    elif len(matches) > 1:
        match_status = 'Rechnungsnr. mehrfach in DATEV'
        assigned = None
    else:
        match_status = 'Kein DATEV-Rechnungstreffer'
        assigned = None
    later_rows.append({
        'source_row': item['source_row'], 'name': name, 'reference': ref,
        'paid': item['paid'], 'date': item['parsed_date'],
        'account': assigned, 'status': match_status,
        'note': item['note'],
    })

accounts = []
for account in sorted(set(datev_accounts) | set(source_accounts)):
    objs = datev_accounts.get(account, [])
    if account in source_accounts:
        name = source_accounts[account]
    else:
        descriptions = [str(o.get('posting_description') or '') for o in objs]
        name = next((x for x in descriptions if x), '')
    accounts.append({'account': account, 'name': name})

split_rows=[]
split_map={}
for r in range(4,30):
    account=split_sheet.cell(r,1).value
    if not isinstance(account,int):
        continue
    ref=str(split_sheet.cell(r,3).value or '')
    part_2024=amount(split_sheet.cell(r,5).value)+amount(split_sheet.cell(r,6).value)
    part_2025=amount(split_sheet.cell(r,8).value)
    item={
        'source_row':r,'account':account,'name':str(split_sheet.cell(r,2).value or ''),
        'ref_2024':ref,'part_2024':part_2024,
        'ref_2025':str(split_sheet.cell(r,7).value or ''),'part_2025':part_2025,
        'note':str(split_sheet.cell(r,9).value or ''),
    }
    split_rows.append(item)
    split_map[ref]=item

for item in later_rows:
    linked=split_map.get(item['reference'])
    item['ref_2025']=linked['ref_2025'] if linked else ''
    item['part_2025']=linked['part_2025'] if linked else None
    if linked:
        if item['account'] is not None and item['account'] != linked['account']:
            item['note']=(item['note']+' ' if item['note'] else '')+f"Kontokonflikt: DATEV {item['account']}, Aufteilung {linked['account']}."
            item['account']=None
            item['status']='Kontokonflikt DATEV/Aufteilung'
        else:
            item['status']='Aufteilung 2024/25: Zahlungszweck prüfen'

source_refs=set(x['reference'] for x in source_rows if x['reference'].startswith('2024-'))
source_refs.update(x['ref_2024'] for x in split_rows)
tracker_rows=[]
tracker=tracker_file['2024']
for r in range(2,tracker.max_row+1):
    raw_ref=tracker.cell(r,7).value
    if not isinstance(raw_ref,(int,float)):
        continue
    ref=f'2024-{int(raw_ref):03d}'
    if ref not in source_refs:
        continue
    evidence=[]
    for c in (10,12,14,16,17,19,20,21,22):
        val=tracker.cell(r,c).value
        if val is not None:
            if isinstance(val,datetime): val=val.strftime('%d.%m.%Y')
            evidence.append(f'{openpyxl.utils.get_column_letter(c)}={val}')
    tracker_rows.append({'source_row':r,'name':str(tracker.cell(r,1).value or ''),
                         'reference':ref,'invoice_amount':amount(tracker.cell(r,9).value),
                         'evidence':'; '.join(evidence)})
tracker_by_ref={item['reference']:item for item in tracker_rows}
for item in split_rows:
    tracker_item=tracker_by_ref.get(item['ref_2024'])
    item['tracker_invoice']=tracker_item['invoice_amount'] if tracker_item else None

storno_rows=[]
storno=tracker_file['Stornorechnungen']
for r in range(2,storno.max_row+1):
    if storno.cell(r,7).value is None and storno.cell(r,1).value is None:
        continue
    dt=storno.cell(r,8).value
    storno_rows.append({'source_row':r,'name':str(storno.cell(r,1).value or ''),
                        'description':str(storno.cell(r,2).value or ''),
                        'reason':str(storno.cell(r,3).value or storno.cell(r,4).value or ''),
                        'number':str(storno.cell(r,7).value or ''),
                        'date':dt.strftime('%Y-%m-%d') if isinstance(dt,datetime) else str(dt or ''),
                        'amount':amount(storno.cell(r,9).value) if isinstance(storno.cell(r,9).value,(int,float)) else None,
                        'note':str(storno.cell(r,10).value or '')})

cases = [
    [10069, '2023-Sammelbetrag', 56387.07, 'Nicht ausbuchen', 'Excel Zeilen 57–100: kein Einzelrechnungsnachweis; DATEV 10069 netto -1.357,00 €.', '2023er Rechnungen einzeln mit DATEV, Zahlungen und Stornos abstimmen.'],
    [10001, '2024-104 Ortner', 883.50, 'Liste bereinigen', 'DATEV 10001: Rechnung und Umbuchung 10069 vom 18.01.2024 bereits ausgeglichen.', 'Posten in Jamso-Liste als ausgeglichen kennzeichnen; keine neue Forderungsausbuchung.'],
    [10105, '2024-214 Doumbia', 3820.00, 'Liste bereinigen', 'DATEV: Storno 2024214ST über 3.820,00 €, Posten bereits ausgeglichen.', 'Stornobeleg zuordnen und Mandantenliste korrigieren.'],
    [10046, '2024-150 Garcia Lopez', 944.00, 'Storno prüfen', 'Stornoliste: Nr. 18258 zu 2024-150 mit -944,00 €; DATEV und Excel zeigen 944,00 € offen.', 'Stornobeleg und DATEV-Erfassung prüfen; danach Gegenbuchung/Auszifferung.'],
    [10049, '2024-185 Annina Benz', 70.00, 'Storno prüfen', 'Stornoliste: Nr. 18252 zu 2024-185 über 70,00 €; DATEV und Excel zeigen 70,00 € offen.', 'Stornobeleg und DATEV-Erfassung prüfen; Quell-Datum 2004 auffällig.'],
    [10059, '2024-165 Devin Searcy', 472.50, 'Storno prüfen', 'Stornoliste nennt Nr. 18253 ohne Betrag; DATEV und Excel zeigen 472,50 € offen.', 'Stornobeleg und tatsächlich wirksamen Betrag klären.'],
    [None, '2024-357/-360/-361', 4370.00, '2025-Stornos prüfen', 'Aufteilungsliste nennt Storno der 2024-Anteile 1.980 € / 1.620 € / 770 € und neue 2025-Rechnungen.', 'Stornobelege und DATEV-Folgebuchungen den drei Rechnungen zuordnen.'],
    [None, '2024-356 Erik Bauer', 2592.00, 'Kontokonflikt prüfen', 'Aufteilung nennt Debitor 10125, DATEV zur Rechnung 2024-356 Debitor 10215; Zahlung entspricht 2025-Rechnung 2.592 €.', 'Debitorennummer und Zahlungszweck 2025 prüfen; nicht auf 2024 offen verrechnen.'],
    [None, '2024-213/-264', 1045.20, 'Rechnungs-Split prüfen', 'Aufteilung E+G liegt um 726,40 € (Lumley) und 318,80 € (Tokarzk) über dem 2024er Tracker-Rechnungsbetrag.', 'Originale 2024- und 2025-Rechnungen samt Storno/Gutschrift abgleichen.'],
    [None, '2024-344/-348', 10.80, 'Betragsabweichung', 'Aufteilung E+G liegt 0,80 € (Jung) und 10,00 € (Schupp) unter dem Tracker-Rechnungsbetrag.', 'Beträge mit Rechnungsoriginalen berichtigen.'],
    [10136, '2024-252 Andrea Friedrich', 720.00, 'Zuordnung prüfen', 'DATEV Andrea +720,00 € offen; Eva Friedrich -720,00 € offen; Excel Zahlung bei Eva.', 'Zahlungsbeleg prüfen, dann debitorisch umbuchen und ausziffern.'],
    [10072, '2024-276 Carla Frey', 1140.00, 'Doppelbuchung prüfen', 'DATEV: ein 1.140-€-Paar ausgeglichen und ein weiterer 1.140-€-Sollposten offen.', 'Rechnung und Buchungssätze prüfen; ggf. Doppelbuchung korrigieren.'],
    [10213, '2024-354 Luis Schley', 1291.00, 'Doppelbuchung prüfen', 'DATEV zwei offene Sollposten zu 1.291,00 €; Excel eine Rechnung und Zahlung 2025.', 'Rechnungsoriginal und 2025-Bankbeleg prüfen; doppelten Satz ggf. stornieren.'],
    [10096, '2024-203 / -223 Lea Wette', 2379.00, 'Differenz klären', 'Excel-Saldo 81,00 €, DATEV-Saldo 2.460,00 €; Zahlung 465 € vs. DATEV 456 €.', 'Rechnungsduplikat, Zusatzrechnung und Zahlungssplit abstimmen.'],
    [10143, '2024-261 Nicole Farrier', 477.00, 'Rückbuchung klären', 'Excel führt zwei Zahlungen zu je 477 €; DATEV zeigt Zahlung und Rückbuchung, netto +477 €.', 'Bankkonto und Rückbuchungsgrund prüfen; Zahlungsstatus korrigieren.'],
    [10161, '2024-282 Brittanne Macey', 345.60, 'Liste ergänzen', 'DATEV: Zahlung 345,60 €; Excel führt nur Rechnung 432,00 €.', 'Zahlung in Mandantenliste ergänzen; Rest 86,40 € bewerten.'],
    [10154, '2024-347 Tkachenko', 680.00, 'Rechnung fehlt', 'DATEV: offene Rechnung 680,00 €; Excel nur Zahlung 544,00 € im Jahr 2025.', 'Rechnung und Zahlung zuordnen, verbleibenden Betrag abstimmen.'],
    [10173, 'Chiara Klein', 1647.00, 'Buchungen abgleichen', 'DATEV zum 31.12.2024 netto 1.647,00 €; Excel enthält Zahlung ohne Rechnungsbetrag 2024-296.', 'Rechnungen 296, 332, 343 und Zahlungen 2025 einzeln zuordnen.'],
    [10177, '2024-300 Ela Centimese', 513.80, 'Rechnung fehlt', 'DATEV 2.569,00 € Rechnung minus 2.055,20 € Zahlung; Excel zeigt nur Zahlung.', 'Rechnung in Liste ergänzen und Rest 513,80 € klären.'],
    [10174, '2024-297 Estelle Gilleron', 764.00, 'Doppelbetrag prüfen', 'DATEV zwei offene Sollbeträge 764,00 € und 550,00 € zur gleichen Rechnungsnr.; Excel 550 €.', 'Rechnungsoriginal und Buchungen prüfen.'],
    [10068, '2024-209 Anna Herzog', 1120.00, '2025-Zahlung ausziffern', 'DATEV 31.12.2024 +1.120,00 €; Excel Zahlung 24.02.2025.', 'Bankbeleg 2025 prüfen und offenen DATEV-Posten ausgleichen.'],
    [10031, '2024-134 Bielefeld', 1360.00, '2025-Zahlung ausziffern', 'DATEV 31.12.2024 +1.360,00 €; Excel Zahlung 02.01.2025.', 'Bankbeleg prüfen und DATEV-Folgejahr ausziffern.'],
    [10166, '2024-288 Korczak', 1726.00, '2025-Zahlung ausziffern', 'DATEV 31.12.2024 +1.726,00 €; Excel Zahlung 03.01.2025.', 'Bankbeleg prüfen und DATEV-Folgejahr ausziffern.'],
    [10178, '2024-301 Botschafter', 1480.00, '2025-Zahlung ausziffern', 'DATEV 31.12.2024 +1.480,00 €; Excel Zahlung 20.01.2025.', 'Bankbeleg prüfen und DATEV-Folgejahr ausziffern.'],
    [10071, '2024-175 / -259 Jara Maier', 6885.00, 'Forderung/Storno prüfen', 'Excel und DATEV je 6.885,00 € netto offen; Stornoliste nennt Nr. 18247 zu 2024-259 ohne Betrag.', 'Storno 259 und Restforderung 175 getrennt belegen; Mahnstand und Werthaltigkeit prüfen.'],
    [10016, '2024-119 Sophie Kegel', 3295.20, 'Forderung prüfen', 'Excel und DATEV je 3.295,20 € netto offen; kein Folgeeingang in Excel.', 'Mahn- und Einziehungsstand belegen; ggf. Wertberichtigung entscheiden.'],
    [None, 'Nachtragszahlungen', 67466.20, 'Einzeln zuordnen', 'Excel Zeilen 201–224 nach Summenzeile; viele Beträge decken 2025-Folgerechnungen aus der Aufteilungsliste.', '2024/25-Rechnung, Debitor und Bankbeleg je Zahlung verbinden; nichts pauschal abziehen.'],
]

payload = {'source_rows': source_rows, 'datev_open_rows': open_rows, 'accounts': accounts, 'later_rows': later_rows, 'split_rows':split_rows,'tracker_rows':tracker_rows,'storno_rows':storno_rows,'cases': cases}
(ROOT / 'workpaper_data.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')

main = [x for x in source_rows if x['section']=='Hauptliste']
total = sum(Decimal(str(x['debit']))-Decimal(str(x['paid'])) for x in main)
cutoff = sum(Decimal(str(x['debit']))-(Decimal(str(x['paid'])) if x['payment_year'] is not None and x['payment_year']<=2024 else Decimal(0)) for x in main)
datev_total = sum(Decimal(str(x['debit']))-Decimal(str(x['credit'])) for x in open_rows)
print(json.dumps({'source_rows':len(source_rows),'datev_open_rows':len(open_rows),'accounts':len(accounts),'later_rows':len(later_rows),'split_rows':len(split_rows),'tracker_rows':len(tracker_rows),'storno_rows':len(storno_rows),'split_account_conflicts':[(s['ref_2024'],s['account'],sorted(account_refs.get(clean_ref(s['ref_2024']),[]))) for s in split_rows if account_refs.get(clean_ref(s['ref_2024'])) and s['account'] not in account_refs[clean_ref(s['ref_2024'])]],'split_amount_conflicts':[(s['ref_2024'],s['part_2024']+s['part_2025'],t['invoice_amount']) for s in split_rows for t in tracker_rows if s['ref_2024']==t['reference'] and abs((s['part_2024']+s['part_2025'])-t['invoice_amount'])>=0.01],'main_total':str(total),'cutoff_total':str(cutoff),'later_total':str(sum(Decimal(str(x['paid'])) for x in later_rows)),'datev_total':str(datev_total)},ensure_ascii=False))
