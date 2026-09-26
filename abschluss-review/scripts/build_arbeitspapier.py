"""Erzeugt das einheitliche Review-Arbeitspapier (Burchardt & Kollegen).

Aufruf:
    python build_arbeitspapier.py <daten.json> <ausgabe.xlsx>

Die Datenstruktur (JSON) ist am Ende dieser Datei dokumentiert (BEISPIEL_JSON).
Das Skript garantiert ein identisches Layout für jeden Review: gleiche Blätter,
Spalten, Farben (Kanzlei-Gold #F7B234), Schweregrad-Formatierung, Dropdowns.
Nach dem Erzeugen zwingend recalc.py des xlsx-Skills laufen lassen.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

GOLD, GREY, INPUT = "F7B234", "F2F2F2", "FFF7E6"
FILL = {
    "A": PatternFill("solid", start_color="FFC7CE"),
    "B": PatternFill("solid", start_color="FFEB9C"),
    "C": PatternFill("solid", start_color="D9D9D9"),
}
SGFONT = {
    "A": Font(color="9C0006", bold=True, size=10),
    "B": Font(color="9C6500", bold=True, size=10),
    "C": Font(bold=True, size=10),
}
HFILL = PatternFill("solid", start_color=GOLD)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F10, F10B = Font(size=10), Font(size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
NUM = '#,##0.00;[Red](#,##0.00);"-"'

RP_COLS = [
    ("Nr.", 8), ("Schweregrad", 12), ("Kategorie", 18), ("Konto / Position", 24),
    ("Betrag GJ (EUR)", 14), ("Betrag VJ (EUR)", 14), ("Feststellung", 52),
    ("Erforderliche Maßnahme", 52), ("Antwort Sachbearbeiter", 36),
    ("Nachweis / Ablage", 20), ("Status", 13), ("Prüfvermerk Partner", 18),
]


def kopf(ws, titel, meta):
    ws["A1"] = "Burchardt & Kollegen — Wirtschaftsprüfer · Steuerberater"
    ws["A1"].font = Font(size=12, bold=True)
    ws["A2"] = "Review-Arbeitspapier: " + titel
    ws["A2"].font = Font(size=11, bold=True)
    ws["A3"] = (f"Mandant: {meta['mandant']} ({meta['mandanten_nr']}) · "
                f"{meta['zeitraum']} · Abschlussart: {meta['abschlussart']}")
    ws["A3"].font = Font(size=9, italic=True)
    ws["A4"] = f"Review erstellt: {meta['review_datum']} (Claude / Skill abschluss-review)"
    ws["A4"].font = Font(size=9, italic=True, color="808080")


def main(datenpfad, ausgabe):
    with open(datenpfad, encoding="utf-8") as f:
        d = json.load(f)
    meta, punkte = d["meta"], d["punkte"]
    wb = Workbook()

    # Deckblatt
    ws = wb.active
    ws.title = "Deckblatt"
    kopf(ws, "Deckblatt", meta)
    zeilen = [
        ("Mandant", meta["mandant"]), ("Mandanten-Nr.", meta["mandanten_nr"]),
        ("Rechtsform / Branche", meta["rechtsform_branche"]),
        ("Abschlussart", meta["abschlussart"]), ("Stichtag / Zeitraum", meta["zeitraum"]),
        ("Verwendete Dokumente", meta["dokumente"]),
        ("Wesentlichkeit", meta["wesentlichkeit"]),
        ("Nichtaufgriffsgrenze", meta["nichtaufgriffsgrenze"]),
        ("Ersteller Review", "Claude (Skill abschluss-review)"),
        ("Sachbearbeiter", ""), ("Abgearbeitet am", ""), ("Freigabe Partner am", ""),
    ]
    r = 6
    for k, v in zeilen:
        ws.cell(r, 1, k).font = F10B
        ws.cell(r, 1).fill = PatternFill("solid", start_color=GREY)
        c = ws.cell(r, 2, v)
        c.font, c.alignment = F10, WRAP
        for col in (1, 2):
            ws.cell(r, col).border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100

    # Reviewpunkte
    ws = wb.create_sheet("Reviewpunkte")
    kopf(ws, "Reviewpunkte", meta)
    HR = 6
    for i, (name, w) in enumerate(RP_COLS, 1):
        c = ws.cell(HR, i, name)
        c.font, c.fill, c.border = Font(bold=True, size=10), HFILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ordnung = {"A": 0, "B": 1, "C": 2}
    punkte = sorted(punkte, key=lambda p: (ordnung[p["schweregrad"]],
                                           -(p.get("betrag_gj") or 0)))
    for j, p in enumerate(punkte):
        r = HR + 1 + j
        vals = [f"R-{j+1:03d}", p["schweregrad"], p["kategorie"], p["konto_position"],
                p.get("betrag_gj"), p.get("betrag_vj"), p["feststellung"],
                p["massnahme"], "", "", "offen", ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border, c.font = BORDER, F10
            c.alignment = WRAP if i in (7, 8, 9) else TOP
            if i in (5, 6):
                c.number_format = NUM
        sg = ws.cell(r, 2)
        sg.fill, sg.font = FILL[p["schweregrad"]], SGFONT[p["schweregrad"]]
        sg.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(r, 9).fill = PatternFill("solid", start_color=INPUT)
        ws.row_dimensions[r].height = 105
    last = HR + len(punkte)
    dv1 = DataValidation(type="list", formula1='"A,B,C"', allow_blank=False)
    dv2 = DataValidation(type="list",
                         formula1='"offen,in Klärung,geklärt,kein Handlungsbedarf"',
                         allow_blank=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"B{HR+1}:B{last}")
    dv2.add(f"K{HR+1}:K{last}")
    ws.auto_filter.ref = f"A{HR}:L{last}"
    ws.freeze_panes = f"A{HR+1}"

    # Summary (Zählung per Formel aus Reviewpunkte-Blatt)
    ws = wb.create_sheet("Summary", 1)
    kopf(ws, "Summary", meta)
    ws["A6"] = "Reviewpunkte nach Schweregrad"
    ws["A6"].font = F10B
    beschr = {"A": "Kritisch — vor Freigabe zwingend klären",
              "B": "Wesentlich — Antwort dokumentieren", "C": "Hinweis"}
    for i, sg in enumerate("ABC"):
        r = 7 + i
        ws.cell(r, 1, sg).font = SGFONT[sg]
        ws.cell(r, 1).fill = FILL[sg]
        ws.cell(r, 2, f'=COUNTIF(Reviewpunkte!B:B,A{r})').font = F10B
        ws.cell(r, 3, beschr[sg]).font = F10
        for col in (1, 2, 3):
            ws.cell(r, col).border = BORDER
    ws.cell(10, 1, "Gesamt").font = F10B
    ws.cell(10, 2, "=SUM(B7:B9)").font = F10B
    for col in (1, 2):
        ws.cell(10, col).border = BORDER
    ws["A12"] = "Top-Punkte"
    ws["A12"].font = F10B
    r = 13
    for t in d["top_punkte"]:
        ws.cell(r, 1, t["nr"]).font = F10B
        c = ws.cell(r, 2, t["text"])
        c.font, c.alignment = F10, WRAP
        for col in (1, 2):
            ws.cell(r, col).border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1
    ws.merge_cells(f"A{r+1}:D{r+1}")
    c = ws.cell(r + 1, 1, "Gesamteinschätzung: " + d["gesamteinschaetzung"])
    c.font, c.alignment = F10, WRAP
    ws.row_dimensions[r + 1].height = 60
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 15

    # Kennzahlen: Basiswerte (blau = Input) + Kennzahlen als Excel-Formeln
    ws = wb.create_sheet("Kennzahlen")
    kopf(ws, "Kennzahlen", meta)
    ws["A6"] = d["kennzahlen"].get("hinweis", "")
    ws["A6"].font = Font(size=9, italic=True)
    r = 8
    for i, t in enumerate(["Basis", "GJ", "VJ"], 1):
        c = ws.cell(r, i, t)
        c.font, c.fill, c.border = F10B, HFILL, BORDER
    b0 = r + 1
    for name, gj, vj in d["kennzahlen"]["basis"]:
        r += 1
        ws.cell(r, 1, name).font = F10
        for col, val in ((2, gj), (3, vj)):
            c = ws.cell(r, col, val)
            c.number_format = NUM
            c.font = Font(size=10, color="0000FF")
        for col in (1, 2, 3):
            ws.cell(r, col).border = BORDER
    r += 2
    for i, t in enumerate(["Kennzahl", "GJ", "VJ", "siehe Punkt"], 1):
        c = ws.cell(r, i, t)
        c.font, c.fill, c.border = F10B, HFILL, BORDER
    for k in d["kennzahlen"]["kennzahlen"]:
        r += 1
        ws.cell(r, 1, k["name"]).font = F10
        # Formeln referenzieren die Basiszeilen: {b0} = erste Basiszeile
        ws.cell(r, 2, k["formel_gj"].format(b=b0)).number_format = k["format"]
        ws.cell(r, 3, k["formel_vj"].format(b=b0)).number_format = k["format"]
        ws.cell(r, 4, k.get("punkt", "—")).font = F10
        for col in (1, 2, 3, 4):
            ws.cell(r, col).border = BORDER
    ws.column_dimensions["A"].width = 32
    for col in "BCD":
        ws.column_dimensions[col].width = 16

    # Buchungsauffälligkeiten (nur wenn vorhanden)
    if d.get("buchungen"):
        ws = wb.create_sheet("Buchungsauffälligkeiten")
        kopf(ws, "Buchungsauffälligkeiten", meta)
        HR = 6
        bcols = [("Konto", 10), ("Datum", 12), ("Betrag (EUR)", 14), ("S/H", 6),
                 ("Gegenkonto", 12), ("Auffälligkeit", 70), ("Reviewpunkt", 12)]
        for i, (name, w) in enumerate(bcols, 1):
            c = ws.cell(HR, i, name)
            c.font, c.fill, c.border = Font(bold=True, size=10), HFILL, BORDER
            ws.column_dimensions[get_column_letter(i)].width = w
        for j, b in enumerate(d["buchungen"]):
            r = HR + 1 + j
            vals = [b["konto"], b["datum"], b["betrag"], b["sh"], b["gegenkonto"],
                    b["auffaelligkeit"], b["punkt"]]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.border, c.font = BORDER, F10
                if i == 3:
                    c.number_format = NUM
                if i == 6:
                    c.alignment = WRAP
            ws.row_dimensions[r].height = 28
        ws.auto_filter.ref = f"A{HR}:G{HR+len(d['buchungen'])}"
        ws.freeze_panes = f"A{HR+1}"

    wb.save(ausgabe)
    print(f"Arbeitspapier gespeichert: {ausgabe} ({len(punkte)} Punkte)")


BEISPIEL_JSON = r"""
{
  "meta": {
    "mandant": "Mustermann GmbH, Dortmund",
    "mandanten_nr": "413885/99999/2025",
    "rechtsform_branche": "GmbH / Handel",
    "abschlussart": "BILANZ (Steuerrecht) — KEINE EÜR",
    "zeitraum": "31.12.2025 · GuV 01.01.–31.12.2025 · Vorjahr 2024",
    "dokumente": "Bilanz+GuV+Kontennachweis (Druck ...); Kontenblätter 01–13/2025",
    "wesentlichkeit": "2.500,00 EUR = 2,5 % des Ergebnisses vor Steuern (100.000,00)",
    "nichtaufgriffsgrenze": "250,00 EUR; qualitative steuerliche Themen unabhängig davon",
    "review_datum": "06.07.2026"
  },
  "punkte": [
    {"schweregrad": "A", "kategorie": "Steuerlich", "konto_position": "4650 Bewirtung",
     "betrag_gj": 5000.0, "betrag_vj": 1200.0,
     "feststellung": "…mit Fundstelle…", "massnahme": "…imperativ, konkret…"}
  ],
  "top_punkte": [{"nr": "R-001", "text": "Einzeiler"}],
  "gesamteinschaetzung": "3–5 Sätze, keine Freigabeaussage.",
  "kennzahlen": {
    "hinweis": "GJ = 12 Monate …",
    "basis": [["Umsatzerlöse", 1000000.0, 950000.0]],
    "kennzahlen": [
      {"name": "Umsatzrendite", "formel_gj": "=B{b}/B{b}", "formel_vj": "=C{b}/C{b}",
       "format": "0.0%", "punkt": "R-001"}
    ]
  },
  "buchungen": [
    {"konto": "4650", "datum": "15.03.2025", "betrag": 2500.0, "sh": "S",
     "gegenkonto": "1200", "auffaelligkeit": "…", "punkt": "R-001"}
  ]
}
"""

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
